"""Isolated, atomic use of the existing desktop Excel import and export adapters.

The worker replaces dialog and connection bindings only in its own process.
No backup/restore entry point is called and no source module is changed.
"""
import json
import os
import sqlite3
import subprocess
import sys
import tempfile
from contextlib import ExitStack, contextmanager, closing
from decimal import Decimal
from pathlib import Path
from unittest.mock import patch

IMPORTS={'suppliers':'import_suppliers','offers':'import_products','purchases':'import_purchases'}
EXPORTS={'suppliers':'export_suppliers','offers':'export_products','purchases':'export_purchases','operating':'export_operating_workbook'}


@contextmanager
def bind_business_connections(conn):
    """Only use in a worker process: all services join one outer transaction."""
    @contextmanager
    def shared(*args,**kwargs):
        yield conn
    with ExitStack() as patches:
        for name,loaded in list(sys.modules.items()):
            if loaded and (name.startswith('services.') or name=='db.connection'):
                for attribute in ('db_transaction','db_read'):
                    if hasattr(loaded,attribute):
                        patches.enter_context(patch.object(loaded,attribute,shared))
        from services import funds_service, historical_receivable_service
        patches.enter_context(patch.object(funds_service,'_connection',shared))
        patches.enter_context(patch.object(historical_receivable_service,'transaction',shared))
        # Project creation owns a connection unless its existing keyword is used.
        # Bind both direct imports and the browser's stored callable to this worker.
        from services import project_service, web_workbench_service
        create_project=project_service.create_project
        def shared_project(data, *, connection=None):
            return create_project(data, connection=conn if connection is None else connection)
        patches.enter_context(patch.object(project_service,'create_project',shared_project))
        patches.enter_context(patch.dict(web_workbench_service.SPECS['projects'],create=shared_project))
        yield


def run(database,operation,module,path,fingerprint=None):
    if module not in (IMPORTS if operation in ('preview','apply') else EXPORTS):
        raise ValueError('未开放此 Excel 操作')
    environment=os.environ.copy()
    environment['SUPPLY_CHAIN_DB_PATH']=str(database)
    environment['PYTHONIOENCODING']='utf-8'
    command=[sys.executable,'-B','-m','web_workbench.transfer','worker',operation,module,str(path),fingerprint or '']
    result=subprocess.run(command,cwd=Path(__file__).resolve().parents[1],env=environment,
                          capture_output=True,text=True,encoding='utf-8',timeout=60)
    try:
        payload=json.loads(result.stdout.splitlines()[-1])
    except (IndexError,ValueError):
        raise ValueError('本机 Excel 操作未完成；未报告成功') from None
    if 'error' in payload:
        raise ValueError(payload['error'])
    return payload


def worker(operation,module,path,expected):
    from openpyxl import load_workbook
    import db.connection as connection
    from services.web_workbench_operations import database_fingerprint
    from pages import import_export_page as legacy
    from services import web_workbench_service as business
    messages=[];errors=[];created=[];row_count=0
    importing=operation in ('preview','apply')
    if importing:
        with closing(load_workbook(path,data_only=False,read_only=True)) as raw,closing(load_workbook(path,data_only=True,read_only=True)) as cached:
            sheet=raw.active
            if sheet.max_row>5001 or sheet.max_column>100:
                raise ValueError('每次最多5000行、100列，请拆分文件')
            for row in sheet.iter_rows(min_row=2):
                if any(cell.value not in (None,'') for cell in row):row_count+=1
                for cell in row:
                    if cell.data_type=='f' and cached.active[cell.coordinate].value is None:
                        raise ValueError(f'{sheet.title}!{cell.coordinate} 公式没有已计算值，请在 Excel 中计算并保存')
            if not row_count:raise ValueError('文件没有可导入的业务行')
    database=Path(connection.DB_PATH)
    if not database.is_file():raise ValueError('明确选择的数据库不存在')
    with tempfile.TemporaryDirectory(prefix='engineering_excel_') as temporary:
        selected=database
        if operation!='apply':
            selected=Path(temporary)/'snapshot.db'
            with closing(sqlite3.connect(f'{database.as_uri()}?mode=ro',uri=True)) as original,closing(sqlite3.connect(selected)) as snapshot:
                original.backup(snapshot)
        conn=sqlite3.connect(selected);conn.row_factory=sqlite3.Row
        try:
            conn.execute('PRAGMA foreign_keys=ON')
            if not importing:
                conn.execute('PRAGMA query_only=ON')
            conn.execute('BEGIN IMMEDIATE' if importing else 'BEGIN')
            if operation=='apply' and database_fingerprint(conn)!=expected:
                raise ValueError('经营资料在预览后已变化，本次 Excel 导入未提交，请重新预览')
            fingerprint=database_fingerprint(conn)
            class Widget:
                def __init__(self,*args,**kwargs):
                    if kwargs.get('text'):errors.append(kwargs['text'])
                def insert(self,*args):errors.append(str(args[-1]))
                def __getattr__(self,name):return lambda *args,**kwargs:None
            def capture(*args,**kwargs):messages.append(str(args[-1]));return True
            target=legacy.ImportExportPage.__new__(legacy.ImportExportPage);target.parent=None
            with bind_business_connections(conn),ExitStack() as patches:
                patches.enter_context(patch.object(connection,'DB_PATH',selected))
                patches.enter_context(patch.object(legacy.filedialog,'askopenfilename',return_value=str(path)))
                patches.enter_context(patch.object(legacy.filedialog,'asksaveasfilename',return_value=str(path)))
                for name in ('showinfo','showwarning','showerror'):patches.enter_context(patch.object(legacy.messagebox,name,capture))
                for name in ('Toplevel','Label','Text'):patches.enter_context(patch.object(legacy.ttk,name,Widget))
                import ui.dialogs as dialogs
                patches.enter_context(patch.object(dialogs,'build_form_dialog',return_value=(Widget(),Widget())))
                patches.enter_context(patch.object(dialogs,'add_form_actions',return_value=None))
                # The legacy product parser otherwise rounds its input through float.
                patches.enter_context(patch.object(legacy,'float',lambda value:Decimal(str(value)),create=True))
                if importing:
                    service=legacy.procurement_service if module=='purchases' else legacy.master_data_service
                    method={'suppliers':'create_supplier','offers':'create_supplier_offer','purchases':'add_purchase_order'}[module]
                    original=getattr(service,method)
                    def tracked(data,*items,**kwargs):
                        if module=='purchases':
                            number=data.get('order_no')
                            if not number:
                                raise ValueError('网页整批导入请补采购单号，不能凭无编号资料确定是否重复业务')
                            if conn.execute('SELECT 1 FROM purchase_orders WHERE order_no=?',(number,)).fetchone():
                                raise ValueError('采购单号已存在：'+number+'；重复导入不会写入')
                        elif module=='suppliers':
                            if any(row['name'].strip()==str(data.get('name','')).strip() for row in business.rows('suppliers')):
                                raise ValueError('供应商名称已存在；请使用档案编辑，重复导入不会写入')
                        elif any(row['supplier_id']==data.get('supplier_id') and row['name']==data.get('name') and row['specification']==data.get('specification','') for row in business.rows('offers')):
                            raise ValueError('相同供应商、材料和规格报价已存在；请编辑报价')
                        record_id=original(data,*items,**kwargs)
                        created.append({'module':module,'id':record_id,'item_count':len(items[0]) if items else 1})
                        return record_id
                    patches.enter_context(patch.object(service,method,tracked))
                    getattr(target,IMPORTS[module])()
                    if errors or sum(row['item_count'] for row in created)!=row_count:
                        raise ValueError('整批未提交：存在失败或被跳过的行。'+ '\n'.join(errors))
                    records=[business.get_record(module,row['id']) for row in created]
                    result={'rows':records,'row_count':row_count,'record_count':len(created),'fingerprint':fingerprint,'messages':messages}
                else:
                    getattr(target,EXPORTS[module])()
                    if not Path(path).is_file():raise ValueError('导出没有生成文件')
                    workbook=load_workbook(path)
                    if module=='purchases':
                        for row in workbook.active.iter_rows(min_row=2):
                            if not row[7].value and row[5].value:
                                supplier=business.partners.get_supplier(int(row[5].value))
                                if supplier:row[7].value=supplier['name']
                    for sheet in workbook:
                        for row in sheet:
                            for cell in row:
                                if cell.data_type=='f':cell.data_type='s'
                    exact=workbook.create_sheet('完整字段与精确分值')
                    exact.append(['台账','记录编号','字段','值（整数分原样保存为文本）'])
                    for key in business.SPECS:
                        for row in business.rows(key):
                            for field,value in business.get_record(key,row['id']).items():
                                text=json.dumps(value,ensure_ascii=False,default=str) if isinstance(value,(dict,list,set)) else str(value) if value is not None else ''
                                exact.append([business.SPECS[key]['title'],str(row['id']),field,text]);exact.cell(exact.max_row,4).data_type='s'
                    workbook.save(path);workbook.close();result={'messages':messages}
            if operation=='apply':conn.commit()
            else:conn.rollback()
            return result
        except Exception:
            conn.rollback();raise
        finally:conn.close()


if __name__=='__main__':
    try:
        _,_,operation,module,path,expected=sys.argv
        print(json.dumps(worker(operation,module,path,expected),ensure_ascii=False,default=str))
    except Exception as error:
        print(json.dumps({'error':str(error)},ensure_ascii=False))
        raise SystemExit(1)
