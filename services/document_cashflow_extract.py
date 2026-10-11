"""Conservative local extraction. No model calls and no inferred dates or totals."""

import hashlib
import json
import re
from datetime import date
from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree

ENGINE_VERSION = 'evidence-understanding-v3'
LIMIT_BYTES = 20 * 1024 * 1024
LIMIT_TEXT = 500_000
KINDS = ('合同正文', '补充协议/附件', '材料清单', '预算', '采购付款安排', '施工计划', '人工补充')


def digest(value):
    if not isinstance(value, bytes):
        value = json.dumps(value, ensure_ascii=False, sort_keys=True).encode('utf-8')
    return hashlib.sha256(value).hexdigest()


def read_document(path):
    source = Path(path).resolve()
    if not source.is_file() or source.stat().st_size > LIMIT_BYTES:
        raise ValueError('请选择存在且不超过20MB的本地资料')
    raw = source.read_bytes()
    if len(raw) > LIMIT_BYTES:
        raise ValueError('请选择不超过20MB的本地资料')
    fragments = []
    recognition = {'ocr_pages': [], 'warnings': []}
    try:
        suffix = source.suffix.lower()
        if suffix in ('.txt', '.md', '.csv'):
            try:
                content = raw.decode('utf-8-sig')
            except UnicodeDecodeError:
                content = raw.decode('gb18030')
            fragments = [{'locator': f'第{i}行', 'text': line}
                         for i, line in enumerate(content.splitlines(), 1) if line.strip()]
        elif suffix == '.pdf':
            from pypdf import PdfReader
            reader = PdfReader(BytesIO(raw))
            if reader.is_encrypted or len(reader.pages) > 200:
                raise ValueError('加密PDF或超过200页的资料请先拆分')
            contents = [page.extract_text() or '' for page in reader.pages]
            scanned = [number for number, content in enumerate(contents, 1) if not content.strip()]
            if scanned:
                from services.document_ocr_service import recognize_pdf_pages
                ocr, recognition['warnings'] = recognize_pdf_pages(raw, scanned)
                recognition['ocr_pages'] = [number for number, text in ocr.items() if text.strip()]
            else:
                ocr = {}
            for page_number, content in enumerate(contents, 1):
                is_ocr = not content.strip()
                if is_ocr:
                    content = ocr.get(page_number, '')
                line_label = 'OCR第' if is_ocr else '第'
                fragments.extend({'locator': f'PDF第{page_number}页/{line_label}{i}行', 'text': line}
                                 for i, line in enumerate(content.splitlines(), 1) if line.strip())
                if not content.strip():
                    fragments.append({'locator': f'PDF第{page_number}页',
                                      'text': '【本页无可读文字，扫描识别未取得文字，待人工补充】'})
        elif suffix == '.docx':
            with ZipFile(BytesIO(raw)) as archive:
                info = archive.getinfo('word/document.xml')
                if info.file_size > LIMIT_BYTES:
                    raise ValueError('Word正文过大，请拆分')
                root = ElementTree.fromstring(archive.read(info))
                ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
                fragments = [{'locator': f'Word第{i}段（含表格，无稳定页码）',
                              'text': ''.join(t.text or '' for t in p.findall('.//w:t', ns))}
                             for i, p in enumerate(root.findall('.//w:p', ns), 1)]
        elif suffix == '.xlsx':
            from openpyxl import load_workbook
            with ZipFile(BytesIO(raw)) as archive:
                if sum(info.file_size for info in archive.infolist()) > 50 * 1024 * 1024:
                    raise ValueError('Excel解压后过大，请拆分')
            workbook = load_workbook(BytesIO(raw), read_only=True, data_only=False)
            try:
                for sheet in workbook:
                    for number, row in enumerate(sheet.iter_rows(values_only=True), 1):
                        if number > 10000:
                            raise ValueError('每张工作表最多10000行，请拆分')
                        fragments.append({'locator': f'工作表「{sheet.title}」第{number}行',
                                          'text': '；'.join(str(v) for v in row if v is not None)})
            finally:
                workbook.close()
        else:
            raise ValueError('支持TXT、MD、CSV、PDF、DOCX和XLSX；独立图片请先转为PDF或人工摘录')
    except ValueError:
        raise
    except Exception as error:
        raise ValueError('资料无法完整读取，请核对格式或改用人工摘录') from error
    fragments = validate_fragments([f for f in fragments if f['text'].strip()])
    return {'name': source.name, 'fragments': fragments, 'source_path': str(source),
            'source_file_sha256': digest(raw), 'recognition': recognition}


def validate_fragments(fragments):
    if not isinstance(fragments, list) or not fragments:
        raise ValueError('资料为空，请填写原文及来源位置')
    result = []
    for fragment in fragments:
        if not isinstance(fragment, dict):
            raise ValueError('每段资料须为带原文与位置的对象')
        locator = str(fragment.get('locator') or '').strip()
        content = str(fragment.get('text') or '').strip()
        if not locator or not content:
            raise ValueError('每段资料必须有原文和条款/页码/行号')
        result.append({'locator': locator, 'text': content})
    if sum(len(f['text']) for f in result) > LIMIT_TEXT:
        raise ValueError('文字超过50万字，请分份关联')
    packets = [fragment['understanding'] for fragment in fragments if 'understanding' in fragment]
    if packets:
        if len(packets) != 1 or 'understanding' not in fragments[0]:
            raise ValueError('资料理解只能与完整原文共同保存一次')
        from services.document_cashflow_semantics import normalize_understanding
        normalize_understanding(packets[0], result)
        result[0]['understanding'] = packets[0]
    return result


def money_minor(value):
    try:
        amount = Decimal(str(value).replace(',', '').replace('，', '').strip())
        if not amount.is_finite() or amount <= 0 or amount * 100 != (amount * 100).to_integral_value():
            raise ValueError
        cents = int(amount * 100)
        if cents > 9_000_000_000_000_000:
            raise ValueError
        return cents
    except (ValueError, InvalidOperation, OverflowError):
        raise ValueError('金额须为大于零、最多两位小数的明确人民币金额') from None


def parse_fragment(fragment, source):
    text = fragment['text']
    if not re.search(r'收款|付款|支付|进度款|质保金|预付款|尾款', text):
        return None
    fields = re.split(r'[；;\n]', text)
    first = fields[0].strip()
    direction = None
    if first.startswith('收款'):
        direction = 'in'
    elif first.startswith('付款'):
        direction = 'out'
    title = re.sub(r'^(收款|付款)\s*[:：]?\s*', '', first).strip() or '未命名节点'
    title = title[:180]
    values = {}
    repeated = set()
    for field in fields[1:]:
        match = re.fullmatch(r'\s*(金额|日期|条件|对方)\s*[:：]\s*(.*?)\s*', field)
        if match:
            label, value = match.groups()
            if label in values:
                repeated.add(label)
            values[label] = value
    issues = []
    amount = None
    due = None
    if not direction:
        issues.append('收支方向待核对：请明确标为收款或付款')
    if '金额' not in repeated:
        match = re.fullmatch(r'[¥￥]?\s*([\d,，]+(?:\.\d{1,2})?)\s*(元|万元)?', values.get('金额', ''))
        if match:
            try:
                factor = 10000 if match.group(2) == '万元' else 1
                amount = money_minor(Decimal(match.group(1).replace(',', '').replace('，', '')) * factor)
            except ValueError:
                pass
    if amount is None:
        issues.append('金额待补充/有多个候选，百分比或公式未自动换算')
    if '日期' not in repeated:
        try:
            due = date.fromisoformat(values.get('日期', '')).isoformat()
            if due != values['日期']:
                due = None
        except (ValueError, KeyError):
            pass
    if due is None:
        issues.append('日期待补充；相对工期、验收或审批日期不自动推定')
    condition = values.get('条件', '')
    if not condition or '条件' in repeated:
        condition = '前置条件待核对'
        issues.append('请核对开票、验收、审批等前置条件')
    if repeated:
        issues.append('重复字段：' + '、'.join(sorted(repeated)))
    return {'key': '', 'title': title, 'direction': direction, 'amount_minor': amount,
            'due_date': due, 'condition': condition, 'counterparty': values.get('对方', ''),
            'issues': issues, 'conflict': False,
            'sources': [{**source, 'locator': fragment['locator'], 'excerpt': text}]}


def extract_rows(documents):
    from services.document_cashflow_semantics import understanding_from_fragments
    groups = {}
    for document in documents:
        source = {key: document[key] for key in ('id', 'name', 'version', 'sha256', 'kind')}
        fragments = json.loads(document['fragments_json'])
        understanding = understanding_from_fragments(fragments)
        if understanding is not None:
            candidates = [{**{key: value for key, value in node.items() if key != 'evidence'},
                           'sources': [{**source, **proof} for proof in node['evidence']]}
                          for node in understanding['nodes']]
        else:
            candidates = [parse_fragment(fragment, source) for fragment in fragments]
        for row in candidates:
            if row:
                groups.setdefault((row['direction'], row['title']), []).append(row)
    result = []
    for group, candidates in groups.items():
        variants = {}
        for row in candidates:
            signature = (row['amount_minor'], row['due_date'], row['condition'], row['counterparty'],
                         row.get('timing', ''))
            if signature in variants:
                variants[signature]['sources'].extend(row['sources'])
                variants[signature]['issues'] = list(dict.fromkeys(variants[signature]['issues'] + row['issues']))
            else:
                variants[signature] = row
        for signature, row in variants.items():
            row['conflict'] = len(variants) > 1
            if row['conflict']:
                row['issues'].append('同名收付节点资料冲突：保留各候选，未计入时间表或计划')
            row['key'] = digest([group, signature])[:24]
            result.append(row)
    return sorted(result, key=lambda row: (row['due_date'] or '9999', row['title'], row['key']))


def timeline(rows, *, opening_minor=0, income_delay_days=0, expense_increase_percent=0):
    from datetime import timedelta
    if type(opening_minor) is not int or opening_minor < 0 or opening_minor > 9_000_000_000_000_000:
        raise ValueError('可用启动资金须为非负整数分')
    if type(income_delay_days) is not int or not 0 <= income_delay_days <= 3650:
        raise ValueError('回款延后天数须为0至3650的整数')
    if type(expense_increase_percent) is not int or not 0 <= expense_increase_percent <= 1000:
        raise ValueError('支出增加比例须为0至1000的整数百分比')
    events = {}
    excluded = []
    for row in rows:
        if row['conflict'] or row['issues'] or row['amount_minor'] is None or not row['due_date'] or not row['direction']:
            excluded.append(row['key'])
            continue
        due = date.fromisoformat(row['due_date'])
        amount = row['amount_minor']
        if row['direction'] == 'in':
            try:
                due += timedelta(days=income_delay_days)
            except OverflowError:
                raise ValueError('情景日期超出支持范围') from None
        else:
            # Round upward to a cent, keeping the estimate conservative and integer-only.
            amount = (amount * (100 + expense_increase_percent) + 99) // 100
        event = events.setdefault(due.isoformat(), {'in_minor': 0, 'out_minor': 0, 'row_keys': []})
        event[row['direction'] + '_minor'] += amount
        event['row_keys'].append(row['key'])
    balance = opening_minor
    minimum = opening_minor
    same_day_minimum = opening_minor
    schedule = []
    for due, event in sorted(events.items()):
        same_day_minimum = min(same_day_minimum, balance - event['out_minor'])
        balance += event['in_minor'] - event['out_minor']
        minimum = min(minimum, balance)
        schedule.append({'date': due, **event, 'balance_minor': balance})
    return {'schedule': schedule, 'gap_minor': max(0, -minimum),
            'outflow_first_gap_minor': max(0, -same_day_minimum), 'excluded_keys': excluded,
            'assumptions': {'opening_minor': opening_minor, 'income_delay_days': income_delay_days,
                            'expense_increase_percent': expense_increase_percent},
            'basis': '仅已知且无冲突节点；假设各节点前置条件按期满足，未知部分未计入，非确定现金流。'}
