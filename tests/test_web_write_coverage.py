"""All registered writes and ordinary forms: cancel, persistence and exact retry."""
import unittest
import base64
import io
from pathlib import Path
from unittest.mock import patch
from urllib.parse import urlencode

from tests import test_web_workbench as web_fixtures
from tests.web_write_cases import operation_fixture, operation_facts, form_fixture, form_facts
from services.web_workbench_operations import OPERATIONS, database_fingerprint
from services.web_workbench_service import schema
from db.connection import db_read


class WriteCoverageTests(unittest.TestCase):
    setUpClass=classmethod(web_fixtures.WebWorkbenchTests.setUpClass.__func__)
    tearDownClass=classmethod(web_fixtures.WebWorkbenchTests.tearDownClass.__func__)
    setUp=web_fixtures.WebWorkbenchTests.setUp
    stop=web_fixtures.WebWorkbenchTests.stop
    request=web_fixtures.WebWorkbenchTests.request

    def snapshot(self):
        with db_read() as conn:
            return database_fingerprint(conn)

    def require_ok(self,path,data=None):
        status,response=self.request(path,data)
        self.assertEqual(status,200,(path,response))
        return response if data is None else response['result']

    def check_operation(self,name):
        values=operation_fixture(name,self.business)
        before=self.snapshot()
        body={'name':name,'record':values}
        canceled=self.require_ok('operation_preview',body)
        self.assertEqual(before,self.snapshot(),'preview writes no business rows')
        self.require_ok('operation_discard',{'token':canceled['token']})
        self.assertEqual(self.request('operation_apply',{'token':canceled['token'],'confirmed':True})[0],400)
        self.assertEqual(before,self.snapshot(),'cancel preserves every table')
        preview=self.require_ok('operation_preview',body)
        self.assertEqual(self.request('operation_apply',{'token':preview['token'],'confirmed':False})[0],400)
        applied=self.require_ok('operation_apply',{'token':preview['token'],'confirmed':True})
        after=self.snapshot()
        with db_read() as conn:
            actual,expected=operation_facts(name,self.business,values,conn)
            self.assertEqual(actual,expected,name)
            self.assertEqual(conn.execute('PRAGMA integrity_check').fetchone()[0],'ok')
            self.assertEqual(conn.execute('PRAGMA foreign_key_check').fetchall(),[])
        self.assertNotEqual(before,after,'actual confirmation leaves persisted evidence')
        repeated=self.require_ok('operation_apply',{'token':preview['token'],'confirmed':True})
        self.assertEqual(repeated,applied)
        self.assertEqual(after,self.snapshot(),'retry changes no row or audit record')

    def check_form(self,module,edit=False):
        fixture=form_fixture(module,self.business,edit=edit)
        record_id=fixture['id']
        version=fixture['record'].get('_version') if edit else None
        parameters={k:v for k,v in {'kind':'save','module':module,'id':record_id,'expected':version}.items() if v is not None}
        before=self.snapshot()
        canceled=self.require_ok('write_ticket?'+urlencode(parameters))
        self.require_ok('write_discard',{'token':canceled['token']})
        body={**fixture,'expected':version,'form_ticket':canceled['token'],'confirmed':True}
        self.assertEqual(self.request('save',body)[0],400)
        self.assertEqual(before,self.snapshot(),'cancel leaves no business changes')
        ticket=self.require_ok('write_ticket?'+urlencode(parameters))
        body['form_ticket']=ticket['token']
        self.assertEqual(self.request('save',{**body,'confirmed':False})[0],400)
        result=self.require_ok('save',body)
        after=self.snapshot()
        actual,expected=form_facts(module,self.business,result)
        self.assertEqual(actual,expected,module)
        if edit:
            row=self.business.get_record(module,result)
            field='description' if module=='construction' else 'basis' if module=='settlements' else 'notes'
            self.assertEqual(row[field],'合成收口修改已保存 '+module)
        self.assertEqual(self.require_ok('save',body),result)
        self.assertEqual(after,self.snapshot(),'exact confirmed retry has no duplicate or audit change')
        altered={**body,'record':{**body['record'],'notes':'不能复用已确认表单改内容'}}
        self.assertEqual(self.request('save',altered)[0],400)
        self.assertEqual(after,self.snapshot())

    def test_save_ticket_absent_canceled_stale_and_restart_are_rejected(self):
        body={'module':'costs','record':{'category':'管理费','amount':'1.01','cost_date':'2026-10-03'},'confirmed':True,'form_ticket':'not-issued'}
        before=self.snapshot()
        self.assertEqual(self.request('save',body)[0],400)
        ticket=self.require_ok('write_ticket?kind=save&module=costs')
        self.business.save_record('plans',{'direction':'out','title':'合成使旧表单失效','amount':'1'})
        current=self.snapshot()
        self.assertEqual(self.request('save',{**body,'form_ticket':ticket['token']})[0],400)
        self.assertEqual(current,self.snapshot())
        # A new service has no old form tickets: retry cannot become a new insert.
        self.server.previews.clear()
        self.assertEqual(self.request('save',{**body,'form_ticket':ticket['token']})[0],400)

    def check_action(self,name):
        b=self.business
        if name=='plan_cancel':values={'id':b.save_record('plans',{'direction':'out','title':'合成取消计划','amount':'100.01'}),'reason':'合成取消依据'}
        elif name=='inspection':values={'id':1,'inspection_status':'已验收','inspection_date':'2026-10-03','inspector':'合成验收人','inspection_notes':'合成验收依据'}
        elif name in ('assign_labor','assign_purchase'):
            values=operation_fixture('governance_labor' if name=='assign_labor' else 'governance_purchase',b)
        elif name=='collection':values={'id':1,'owner_name':'合成直接跟进','due_date':'2026-10-03','status':'pending'}
        elif name=='project_close':values=operation_fixture(name,b)
        elif name=='project_site':values={'id':1,'project_id':1,'site_name':'合成新增地点','site_code':'CLOSE-EXTRA-SITE','address':'合成地址'}
        else:
            values=operation_fixture(name,b)
            values.update(project_id=values['id'],request_key='CLOSE-DIRECT-HISTORY')
        parameters={'kind':'action','action':name}
        if values.get('id'):parameters['id']=values['id']
        before=self.snapshot()
        canceled=self.require_ok('write_ticket?'+urlencode(parameters))
        body={'action':name,'record':values,'form_ticket':canceled['token'],'confirmed':True}
        self.require_ok('write_discard',{'token':canceled['token']})
        self.assertEqual(self.request('action',body)[0],400)
        self.assertEqual(before,self.snapshot())
        body['form_ticket']=self.require_ok('write_ticket?'+urlencode(parameters))['token']
        self.assertEqual(self.request('action',{**body,'confirmed':False})[0],400)
        result=self.require_ok('action',body)
        after=self.snapshot()
        with db_read() as conn:
            if name=='plan_cancel':self.assertEqual(conn.execute('SELECT status FROM fund_plans WHERE id=?',(values['id'],)).fetchone()[0],'cancelled')
            elif name=='inspection':self.assertEqual(tuple(conn.execute('SELECT inspection_status,inspection_date,inspector FROM construction_records WHERE id=1').fetchone()),('已验收','2026-10-03','合成验收人'))
            elif name=='collection':self.assertEqual(conn.execute('SELECT owner_name FROM collection_cases WHERE project_id=1').fetchone()[0],'合成直接跟进')
            elif name=='project_site':self.assertEqual(conn.execute('SELECT name FROM project_sites WHERE id=?',(result,)).fetchone()[0],'合成新增地点')
            else:
                actual,expected=operation_facts({'assign_labor':'governance_labor','assign_purchase':'governance_purchase'}.get(name,name),b,values,conn)
                self.assertEqual(actual,expected)
        self.assertNotEqual(before,after)
        self.assertEqual(self.require_ok('action',body),result)
        self.assertEqual(after,self.snapshot())

    def check_csv(self,module):
        csvs={'projects':'name,project_code\n合成CSV项目甲,CLOSE-CSV-A\n合成CSV项目乙,CLOSE-CSV-B\n',
            'suppliers':'name,contact,contact_phone\n合成CSV供应商甲,合成联系人甲,13800000001\n合成CSV供应商乙,合成联系人乙,13800000002\n',
            'customers':'legal_name,credit_limit\n合成CSV客户甲,100.01\n合成CSV客户乙,100.02\n',
            'offers':'supplier_id,name,specification,unit,price,tax_rate_percent\n2,合成CSV材料甲,S1,台,100.01,13\n2,合成CSV材料乙,S2,台,100.02,13\n'}
        upload={'module':module,'name':'synthetic.csv','content':base64.b64encode(csvs[module].encode('utf-8-sig')).decode()}
        before=self.snapshot()
        preview=self.require_ok('csv_preview',upload)
        self.assertEqual(before,self.snapshot(),'preview must remain inside clone, including project auto-sites')
        self.require_ok('csv_discard',{'token':preview['token']})
        self.assertEqual(self.request('csv_apply',{'token':preview['token'],'confirmed':True})[0],400)
        self.assertEqual(before,self.snapshot())
        preview=self.require_ok('csv_preview',upload)
        body={'token':preview['token'],'confirmed':True}
        self.assertEqual(self.request('csv_apply',{**body,'confirmed':False})[0],400)
        result=self.require_ok('csv_apply',body)
        self.assertEqual(result['count'],2)
        rows=[self.business.get_record(module,row['id']) for row in result['rows']]
        if module=='projects':self.assertEqual([row['project_code'] for row in rows],['CLOSE-CSV-A','CLOSE-CSV-B'])
        elif module=='suppliers':self.assertEqual([row['contact_phone'] for row in rows],['13800000001','13800000002'])
        elif module=='customers':self.assertEqual([row['credit_limit_minor'] for row in rows],[10001,10002])
        else:self.assertEqual([row['quoted_price_minor'] for row in rows],[10001,10002])
        after=self.snapshot()
        self.assertEqual(self.require_ok('csv_apply',body),result)
        self.assertEqual(after,self.snapshot())
        self.assertEqual(self.request('csv_preview',upload)[0],400)
        self.assertEqual(after,self.snapshot())

    def check_upload(self,photos=False):
        from services.attachment_service import attachment_path
        if photos:
            from PIL import Image
            file=io.BytesIO();Image.new('RGB',(10,10),'blue').save(file,format='PNG')
            raw=file.getvalue();kind='photos_upload'
            parameters={'kind':kind,'id':1,'expected':self.business.get_record('construction',1)['_version']}
            body={'id':1,'expected':parameters['expected'],'category':'施工现场','notes':'合成照片来源','files':[{'name':'synthetic.png','content':base64.b64encode(raw).decode()}]}
        else:
            raw='合成普通附件条款：款项待确认'.encode();kind='upload_attachment'
            parameters={'kind':kind,'entity':'contract','id':1}
            body={'entity':'contract','id':1,'category':'合同正文','description':'合成附件来源','name':'synthetic.txt','content':base64.b64encode(raw).decode()}
        before=self.snapshot()
        canceled=self.require_ok('write_ticket?'+urlencode(parameters))
        body.update(form_ticket=canceled['token'],confirmed=True)
        self.require_ok('write_discard',{'token':canceled['token']})
        self.assertEqual(self.request(kind,body)[0],400)
        self.assertEqual(before,self.snapshot())
        body['form_ticket']=self.require_ok('write_ticket?'+urlencode(parameters))['token']
        self.assertEqual(self.request(kind,{**body,'confirmed':False})[0],400)
        result=self.require_ok(kind,body)
        after=self.snapshot()
        with db_read() as conn:
            row=conn.execute('SELECT * FROM business_attachments ORDER BY id DESC LIMIT 1').fetchone()
            self.assertEqual(row['original_name'],'synthetic.png' if photos else 'synthetic.txt')
            stored=attachment_path(row['file_path']);self.assertEqual(stored.read_bytes(),raw)
            if photos:self.assertEqual(conn.execute('SELECT photo_type FROM construction_photos WHERE id=?',(result[0],)).fetchone()[0],'施工现场')
        self.assertEqual(self.require_ok(kind,body),result)
        self.assertEqual(after,self.snapshot())
        self.assertEqual(stored.read_bytes(),raw)
        self.assertEqual(self.request(kind,{**body,'category':'不能用旧令牌改分类'})[0],400)
        self.assertEqual(after,self.snapshot())

    def test_supplier_contact_details_require_named_contact(self):
        before=self.snapshot()
        status,response=self.request('save',{'module':'suppliers','record':{'name':'合成不丢联系电话','contact_phone':'13800000000'},'confirmed':True})
        self.assertEqual(status,400,response)
        self.assertIn('联系人姓名',response['error'])
        self.assertEqual(before,self.snapshot())

    def test_attachment_upload_cancel_save_and_retry(self):self.check_upload()
    def test_photo_upload_cancel_save_and_retry(self):self.check_upload(photos=True)

    def test_attachment_staging_does_not_require_same_volume(self):
        from web_workbench import extended
        from services.attachment_service import attachment_path
        source = Path(self.temp.name) / 'staging-source.txt'
        source.write_bytes(b'original attachment')
        payload = {'entity':'contract','id':1,'files':[{'path':str(source),'name':'original.txt'}],
                   'fingerprint':self.snapshot()}
        with patch.object(Path,'replace',side_effect=OSError(18,'cross-volume rename is forbidden')):
            result = extended.worker('attachment_apply',payload)
        with db_read() as conn:
            row = conn.execute('SELECT * FROM business_attachments WHERE id=?',(result,)).fetchone()
        self.assertEqual(row['original_name'],'original.txt')
        self.assertEqual(attachment_path(row['file_path']).read_bytes(),b'original attachment')
        self.assertEqual(source.read_bytes(),b'original attachment')

    def check_excel(self,module):
        from openpyxl import Workbook,load_workbook
        from contextlib import closing
        from pathlib import Path
        from web_workbench import transfer
        file=Path(self.temp.name)/('close-'+module+'.xlsx')
        if module=='offers':
            workbook=Workbook();sheet=workbook.active
            sheet.append(['ID','供应商ID','供应商名称','产品名称','规格','单位','材料单价（未税）','税率（%）','含税单价','备注','报价口径'])
            sheet.append([None,2,'合成供应商','合成Excel收口报价','S1','台','100.01',13,None,'合成批次来源','exclusive'])
            workbook.save(file);workbook.close()
        else:
            transfer.run(self.path,'export','purchases',file)
            with closing(load_workbook(file)) as workbook:
                sheet=workbook.active
                row=list(next(sheet.iter_rows(min_row=2,max_row=2,values_only=True)))
                row[0]=None;row[1]='CLOSE-EXCEL-PO';row[18]='100.01'
                for index,value in enumerate(row,start=1):sheet.cell(2,index).value=value
                second=list(row);second[18]='0';sheet.append(second);workbook.save(file)
        upload={'module':module,'name':file.name,'content':base64.b64encode(file.read_bytes()).decode()}
        before=self.snapshot()
        preview=self.require_ok('excel_preview',upload)
        self.assertEqual(before,self.snapshot())
        self.require_ok('excel_discard',{'token':preview['token']})
        self.assertEqual(self.request('excel_apply',{'token':preview['token'],'confirmed':True})[0],400)
        self.assertEqual(before,self.snapshot())
        preview=self.require_ok('excel_preview',upload)
        body={'token':preview['token'],'confirmed':True}
        self.assertEqual(self.request('excel_apply',{**body,'confirmed':False})[0],400)
        result=self.require_ok('excel_apply',body)
        row=self.business.get_record(module,result['rows'][0]['id'])
        if module=='offers':self.assertEqual(row['quoted_price_minor'],10001)
        else:
            self.assertEqual(row['order_no'],'CLOSE-EXCEL-PO')
            self.assertEqual(len(row['items']),2)
            self.assertEqual(row['freight_amount_cents'],10001)
            self.assertEqual(row['total_amount_cents'],2*10170000+10001)
        after=self.snapshot()
        self.assertEqual(self.require_ok('excel_apply',body),result)
        self.assertEqual(after,self.snapshot())
        self.assertEqual(self.request('excel_preview',upload)[0],400)
        self.assertEqual(after,self.snapshot())

    def test_excel_offer_cancel_save_and_retry(self):self.check_excel('offers')
    def test_excel_purchase_cancel_save_and_retry(self):self.check_excel('purchases')

    def test_manual_input_pdf_attach_cancel_stale_source_and_retry(self):
        import hashlib
        from pathlib import Path
        from tests.test_web_workbench_extended import synthetic_invoice_pdf
        from services.attachment_service import attachment_path
        values=form_fixture('input_invoices',self.business)['record']
        values.update(amount='113',net_amount='100',tax_amount='13')
        record_id=self.business.save_record('input_invoices',values)
        file=Path(self.temp.name)/'close-input-original.pdf'
        synthetic_invoice_pdf(file,number=values['invoice_no'],buyer=values['buyer_name'],seller=values['seller_name'])
        upload={'name':file.name,'content':base64.b64encode(file.read_bytes()).decode()}
        before=self.snapshot()
        preview=self.require_ok('input_pdf_preview',upload)
        body={'token':preview['token'],'id':record_id,'expected':self.business.get_record('input_invoices',record_id)['_version'],'confirmed':True}
        self.require_ok('cashflow/discard',{'token':preview['token']})
        self.assertEqual(self.request('input_pdf_attach',body)[0],400)
        self.assertEqual(before,self.snapshot())
        preview=self.require_ok('input_pdf_preview',upload)
        body['token']=preview['token']
        self.business.save_record('input_invoices',{'notes':'合成旧票版本变化'},record_id,body['expected'])
        current=self.snapshot()
        self.assertEqual(self.request('input_pdf_attach',body)[0],400)
        self.assertEqual(current,self.snapshot())
        preview=self.require_ok('input_pdf_preview',upload)
        body.update(token=preview['token'],expected=self.business.get_record('input_invoices',record_id)['_version'])
        self.assertEqual(self.request('input_pdf_attach',{**body,'confirmed':False})[0],400)
        result=self.require_ok('input_pdf_attach',body)
        after=self.snapshot()
        row=self.business.get_record('input_invoices',record_id)
        self.assertEqual(row['gross_minor'],11300)
        self.assertEqual(row['file_sha256'],hashlib.sha256(file.read_bytes()).hexdigest())
        self.assertEqual(attachment_path(row['file_path']).read_bytes(),file.read_bytes())
        self.assertEqual(self.require_ok('input_pdf_attach',body),result)
        self.assertEqual(after,self.snapshot())
        self.assertEqual(self.request('input_pdf_attach',{**body,'expected':'different-version'})[0],400)
        self.assertEqual(after,self.snapshot())


def operation_test(name):
    def run(self):self.check_operation(name)
    return run


def form_test(module,edit):
    def run(self):self.check_form(module,edit)
    return run


for name,definition in OPERATIONS.items():
    if not definition['read']:
        setattr(WriteCoverageTests,'test_operation_'+name,operation_test(name))
for module,definition in schema().items():
    for edit,capability in ((False,'can_create'),(True,'can_update')):
        if definition[capability]:
            setattr(WriteCoverageTests,'test_'+('edit_' if edit else 'create_')+module,form_test(module,edit))

for action in ('plan_cancel','inspection','assign_labor','assign_purchase','collection','project_close','project_site','historical_receipt'):
    setattr(WriteCoverageTests,'test_action_'+action,lambda self,name=action:self.check_action(name))
for module in ('projects','suppliers','customers','offers'):
    setattr(WriteCoverageTests,'test_csv_'+module,lambda self,name=module:self.check_csv(name))


if __name__=='__main__':unittest.main()
