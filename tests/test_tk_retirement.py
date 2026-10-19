"""The active workbench must start and transfer Excel without the retired UI."""
import ast
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest


ROOT = Path(__file__).resolve().parents[1]


class TkRetirementTests(unittest.TestCase):
    def test_retired_sources_are_not_part_of_the_runtime(self):
        for name in ('main.py', 'run.bat', 'pages', 'ui'):
            with self.subTest(name=name):
                self.assertFalse((ROOT / name).exists())

    def test_production_imports_do_not_depend_on_the_retired_framework(self):
        banned = {'pages', 'ui', 'tkinter', 'ttkbootstrap', 'main'}
        paths = [ROOT / 'ai_engine.py', ROOT / 'ai_client.py']
        for folder in ('services', 'db', 'web_workbench'):
            paths.extend((ROOT / folder).rglob('*.py'))
        for path in paths:
            with self.subTest(path=str(path.relative_to(ROOT))):
                tree = ast.parse(path.read_text(encoding='utf-8-sig'))
                for node in ast.walk(tree):
                    if isinstance(node, ast.Import):
                        modules = [name.name for name in node.names]
                    elif isinstance(node, ast.ImportFrom) and not node.level:
                        modules = [node.module or '']
                    else:
                        continue
                    self.assertFalse(banned & {name.split('.')[0] for name in modules})

    def test_dependencies_are_unique_and_do_not_require_tk(self):
        packages = [line.split('==')[0].casefold() for line in
                    (ROOT / 'requirements.txt').read_text(encoding='utf-8').splitlines()
                    if line.strip() and not line.startswith('#')]
        self.assertEqual(len(packages), len(set(packages)))
        self.assertNotIn('ttkbootstrap', packages)

    def test_fresh_process_exports_all_workbooks_with_ui_imports_blocked(self):
        script = '''
import importlib.abc, logging, sys
from pathlib import Path
class NoOldUI(importlib.abc.MetaPathFinder):
    def find_spec(self, fullname, path=None, target=None):
        if fullname.split('.')[0] in {'pages','ui','tkinter','ttkbootstrap'}:
            raise ImportError('Retired UI dependency: ' + fullname)
sys.meta_path.insert(0, NoOldUI())
logging.disable(logging.CRITICAL)
from db.schema import init_db
init_db()
from web_workbench.synthetic import seed
seed()
from web_workbench import server, desktop, assistant, backup, extended, transfer
from openpyxl import load_workbook
for module in transfer.EXPORTS:
    path = Path(sys.argv[1]) / (module + '.xlsx')
    transfer.worker('export', module, path, '')
    workbook = load_workbook(path)
    assert '完整字段与精确分值' in workbook.sheetnames
    workbook.close()
assert not any(name.split('.')[0] in {'pages','ui','tkinter','ttkbootstrap'} for name in sys.modules)
print('Headless startup and four Excel exports passed')
'''
        with tempfile.TemporaryDirectory(prefix='no_tk_workbench_') as directory:
            environment = os.environ.copy()
            environment.update(SUPPLY_CHAIN_DB_PATH=str(Path(directory) / 'synthetic.db'),
                               SUPPLY_CHAIN_ATTACHMENTS_PATH=str(Path(directory) / 'attachments'),
                               PYTHONIOENCODING='utf-8')
            result = subprocess.run([sys.executable, '-B', '-X', 'utf8', '-c', script, directory],
                                    cwd=ROOT, env=environment, capture_output=True, text=True,
                                    encoding='utf-8', timeout=60)
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            self.assertIn('four Excel exports passed', result.stdout)


if __name__ == '__main__':
    unittest.main()
