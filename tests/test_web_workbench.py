"""Local browser adapter and HTTP boundaries, using synthetic databases only."""
import base64
import json
import sqlite3
import tempfile
import threading
import unittest
import urllib.request
import urllib.error
from contextlib import closing
from pathlib import Path
from unittest.mock import patch


class WebWorkbenchTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        import db.connection as connection
        import db.migration_runner as runner
        from db.schema import init_db
        from web_workbench.synthetic import seed
        cls.temp = tempfile.TemporaryDirectory(prefix='web_synthetic_')
        cls.template = Path(cls.temp.name) / 'template.db'
        with patch.object(connection,'DB_PATH',cls.template), patch.object(runner,'DB_PATH',cls.template):
            init_db()
            seed()

    @classmethod
    def tearDownClass(cls):
        cls.temp.cleanup()

    def setUp(self):
        import db.connection as connection
        from services import web_workbench_service as business
        from web_workbench.server import WorkbenchServer
        self.path = Path(self.temp.name) / (self.id().split('.')[-1] + '.db')
        with closing(sqlite3.connect(self.template)) as source, closing(sqlite3.connect(self.path)) as target:
            source.backup(target)
        self.db_patch = patch.object(connection,'DB_PATH',self.path)
        self.db_patch.start()
        self.addCleanup(self.db_patch.stop)
        self.business = business
        self.server = WorkbenchServer(0,read_only=False,demo=True)
        self.thread = threading.Thread(target=self.server.serve_forever,daemon=True)
        self.thread.start()
        self.addCleanup(self.stop)

    def stop(self):
        self.server.shutdown()
        self.server.server_close()
        self.thread.join(timeout=2)

    def test_navigation_and_scrollbar_assets_are_served(self):
        with urllib.request.urlopen(self.server.origin + '/') as response:
            html = response.read().decode('utf-8')
        self.assertIn('id="navigation-back"', html)
        for filename, symbol in (
            ('navigation.js', 'WorkbenchNavigation'),
            ('table-scrollbar.js', 'WorkbenchTableScrollbar'),
        ):
            with self.subTest(filename=filename):
                self.assertIn(f'src="/{filename}"', html)
                with urllib.request.urlopen(self.server.origin + '/' + filename) as response:
                    self.assertEqual(response.status, 200)
                    self.assertIn(symbol, response.read().decode('utf-8'))

    def request(self,path,data=None,*,headers=None):
        if path in ('save','action') and data and not data.get('form_ticket'):
            from urllib.parse import urlencode
            values={k:v for k,v in {'kind':path,'module':data.get('module'),'action':data.get('action'),'id':data.get('id') or data.get('record',{}).get('id'),'expected':data.get('expected')}.items() if v is not None}
            status,ticket=self.request('write_ticket?'+urlencode(values))
            if status==200:data={**data,'form_ticket':ticket['token']}
        request_headers = {'Origin':self.server.origin,'X-Workbench-Token':self.server.token,
                           'Content-Type':'application/json'}
        request_headers.update(headers or {})
        request = urllib.request.Request(self.server.origin+'/api/'+path,
            data=json.dumps(data).encode() if data is not None else None,headers=request_headers)
        try:
            response=urllib.request.urlopen(request)
        except urllib.error.HTTPError as error:
            response=error
        with response:
            return response.status,json.load(response)

    def test_readonly_assistant_and_preview_never_write_business_or_conversations(self):
        import time
        from services import ai_query_tools
        self.server.read_only = True
        before = self.path.read_bytes()
        plan = {'queries': [{'dataset': 'collections', 'metrics': ['amount_minor'], 'result_mode': 'summary', 'end_date':'2026-10-04'}]}
        self.server.assistant._engine = lambda question, **kwargs: ai_query_tools.execute_plan(plan, db_path=self.path)
        status, created = self.request('assistant/new', {})
        self.assertEqual(status, 200, created)
        identity = created['result']['session_id']
        status, _ = self.request('assistant/ask', {'session_id': identity, 'question': '全部实际回款'})
        self.assertEqual(status, 200)
        for _ in range(100):
            status, snapshot = self.request('assistant/state?session_id=' + identity)
            if not snapshot['busy']:
                break
            time.sleep(.01)
        self.assertFalse(snapshot['busy'])
        self.assertIn('sources', snapshot['messages'][-1])
        self.assertTrue(snapshot['messages'][-1]['sources'])
        self.assertEqual(self.request('assistant/state?session_id=' + identity, headers={'X-Workbench-Token':'wrong'})[0], 400)
        self.assertEqual(self.request('assistant/new', {}, headers={'Origin':'https://example.invalid'})[0], 403)
        self.assertEqual(self.request('assistant/unknown', {})[0], 403)
        self.assertEqual(self.request('assistant/scope', {'session_id':identity,'project_id':999999})[0], 400)
        self.assertEqual(self.request('save', {'module':'projects','record':{'name':'禁止保存'}})[0], 403)
        self.assertEqual(self.path.read_bytes(), before)

    def test_purchase_preview_same_decimal_rules_and_no_writes(self):
        self.server.read_only = True
        before = self.path.read_bytes()
        item = {'quantity':'996', 'unit_price':'1.89', 'price_basis':'inclusive', 'tax_percent':'13'}
        status, response = self.request('purchase_preview', {'items':[item], 'freight':'1.01'})
        self.assertEqual(status, 200, response)
        self.assertEqual(response['result']['items'][0]['line_amount_cents'], '188244')
        self.assertEqual(response['result']['total_amount_cents'], '188345')
        weight = {'settlement_mode':'weight','net_weight':'1.2345','weight_unit':'吨','weight_unit_price':'3250.1234',
                  'price_basis':'inclusive','tax_percent':'13'}
        total = {'settlement_mode':'total','settlement_total':'123.45','price_basis':'exclusive','tax_percent':'13'}
        status, response = self.request('purchase_preview', {'items':[weight,total]})
        self.assertEqual(status, 200, response)
        self.assertEqual(response['result']['items'][0]['line_amount_cents'], '401228')
        self.assertEqual(response['result']['items'][1]['line_amount_cents'], '13950')
        for invalid in ([], [dict(item, quantity='NaN')], [dict(item, tax_percent='101')]):
            self.assertEqual(self.request('purchase_preview', {'items':invalid})[0], 400)
        self.assertEqual(self.path.read_bytes(), before)

    def test_entry_options_keep_completed_for_edit_not_new_work_or_cost(self):
        original = self.business.get_record('projects', 1)
        self.business.projects.update_project(1, {**original, 'status':'已完工'})
        choices = self.business.options()
        self.assertTrue(any(item['value'] == 1 for item in choices['projects']))
        self.assertFalse(any(item['value'] == 1 for item in choices['active_projects']))

    def test_quote_options_include_specification_supplier_unit_price_and_current_records(self):
        from services.web_workbench_operations import database_fingerprint
        from db.connection import db_read
        first = self.business.save_record('offers', {
            'supplier_id': 2, 'name': '合成H型钢', 'specification': '248×124',
            'unit': '支', 'price': '1050.00', 'price_basis': 'inclusive', 'tax_rate_percent': '9',
        })
        second = self.business.save_record('offers', {
            'supplier_id': 2, 'name': '合成H型钢', 'specification': '198×99',
            'unit': '支', 'price': '760.00', 'price_basis': 'exclusive', 'tax_rate_percent': '9',
        })
        with db_read() as conn:
            before = database_fingerprint(conn)
        status, result = self.request('options')
        self.assertEqual(status, 200)
        options = {int(row['value']): row['label'] for row in result['offers']}
        suppliers = {int(row['value']): row['supplier_id'] for row in result['offers']}
        for identity, spec, price in ((first, '248×124', '含税 ¥1050.00/支'),
                                      (second, '198×99', '未税 ¥760.00/支')):
            quote = self.business.get_record('offers', identity)
            self.assertEqual(suppliers[identity],quote['supplier_id'])
            self.assertIn(spec, options[identity])
            self.assertIn(price, options[identity])
            self.assertIn(quote['name'], options[identity])
            self.assertIn(f'报价#{identity}', options[identity])
        self.assertNotEqual(options[first], options[second])
        with db_read() as conn:
            self.assertEqual(database_fingerprint(conn), before)

    def test_date_defaults_refresh_each_day_and_templates_do_not_change(self):
        from datetime import date
        original = self.business.SPECS['costs']['fields']
        original_default = next(row['default'] for row in original if row['name'] == 'cost_date')
        with patch.object(self.business, 'date') as clock:
            clock.today.return_value = date(2026, 10, 12)
            clock.fromisoformat.side_effect = date.fromisoformat
            schema = self.business.schema()
            for spec in schema.values():
                for row in spec['fields']:
                    if row.get('default_today'):
                        self.assertEqual(row['default'], '2026-10-12')
            identity = self.business.save_record('costs', {
                'project_id': 1, 'category': '工程施工 / 材料辅料', 'amount': '1.00',
            })
            self.assertEqual(self.business.get_record('costs', identity)['cost_date'], '2026-10-12')
            clock.today.return_value = date(2026, 10, 13)
            refreshed = self.business.schema()
            self.assertEqual(next(row['default'] for row in refreshed['costs']['fields']
                                  if row['name'] == 'cost_date'), '2026-10-13')
        self.assertEqual(next(row['default'] for row in original if row['name'] == 'cost_date'), original_default)

    def test_purchase_choices_preserve_legacy_facts_and_reject_new_unknown_statuses(self):
        fields = {row['name']: row for row in self.business.schema()['purchases']['fields']}
        for key, facts in {
            'invoice_status': ['有发票', '收据', '无发票', '未确认'],
            'payment_status': ['已付款', '未付款', '未确认'],
            'payment_method': ['现金', '对公转账', '员工垫付', '未记录'],
        }.items():
            self.assertTrue(set(facts).issubset(fields[key]['choices']), key)
        with self.assertRaisesRegex(ValueError, '付款状态.*无效'):
            self.business.save_record('purchases', {
                'project_id': 1, 'purchase_date': '2026-10-12', 'payment_status': '合成无效状态',
            })

    def test_purchase_rows_show_all_original_material_names_and_specifications(self):
        identity = self.business.save_record('purchases', {
            'project_id': 1, 'supplier_id': 2, 'purchase_date': '2026-10-13',
            'items': [
                {'material_name_snapshot': '合成割枪', 'specification_snapshot': 'G01-100',
                 'quantity': '1', 'unit_snapshot': '把', 'unit_price': '68', 'price_basis': 'inclusive'},
                {'material_name_snapshot': '合成气割接头', 'specification_snapshot': 'M12',
                 'quantity': '5', 'unit_snapshot': '只', 'unit_price': '3', 'price_basis': 'inclusive'},
                {'material_name_snapshot': '合成调压阀', 'quantity': '1',
                 'unit_snapshot': '只', 'unit_price': '64', 'price_basis': 'inclusive'},
            ],
        })
        before = self.business.get_record('purchases', identity)
        for keyword in ('', '割枪', 'M12', '调压阀'):
            status, response = self.request('rows?module=purchases&project_id=1&keyword=' + urllib.parse.quote(keyword))
            self.assertEqual(status, 200)
            row = next(row for row in response['rows'] if row['id'] == identity)
            self.assertEqual(row['material_details'], '合成割枪 · G01-100\n合成气割接头 · M12\n合成调压阀')
            self.assertEqual(row['total_amount_cents'], '14700')
            self.assertEqual(row['item_count'], 3)
        self.assertEqual(self.business.get_record('purchases', identity), before)
        columns = self.business.schema()['purchases']['columns']
        self.assertEqual(columns[1], ('material_details', '材料名称 / 规格'))
        self.assertEqual(columns[-1][0], 'order_no')

    def test_old_exclusive_quotes_without_explicit_quoted_price_still_open(self):
        from db.connection import db_transaction
        # Simulate a real pre-migration offer in the isolated fixture, not the formal ledger.
        original = self.business.get_record('offers', 1)
        with db_transaction() as conn:
            conn.execute("UPDATE supplier_offers SET quoted_price_minor=NULL, price_basis='exclusive' WHERE id=1")
        status, choices = self.request('options')
        self.assertEqual(status, 200, choices)
        option = next(row for row in choices['offers'] if int(row['value']) == 1)
        self.assertIn('未税 ¥4500.00/吨', option['label'])
        status, quote = self.request('record?module=offers&id=1')
        self.assertEqual(status, 200, quote)
        self.assertEqual(quote['price'], original['price'])

    def test_attention_and_entry_endpoints_reject_read_only_writes(self):
        self.server.read_only=True
        status,attention=self.request('attention')
        self.assertEqual(status,200);self.assertEqual(len(attention['checks']),2)
        for endpoint,payload in (
            ('attention/day_open',{}),
            ('attention/check',{'kind':'worklogs','day':attention['day'],'status':'done','basis':attention['checks'][0]['basis'],'confirmed':True}),
            ('entry/preview',{'records':[],'request_key':'synthetic-readonly'}),
            ('entry/apply',{'token':'unknown','confirmed':True}),
        ):
            with self.subTest(endpoint=endpoint):self.assertEqual(self.request(endpoint,payload)[0],403)

    def test_http_entry_approval_and_duplicate_request_boundary(self):
        payload={'request_key':'synthetic-http-entry-once','records':[{'module':'costs','data':{
            'project':'厂房钢结构','date':'2026-10-13','category':'工程施工 / 材料辅料',
            'description':'合成工具','amount':'25','paid':True}}]}
        status,preview=self.request('entry/preview',payload)
        self.assertEqual(status,200);draft=preview['result'];self.assertTrue(draft['can_apply'])
        self.assertEqual(self.request('entry/apply',{'token':draft['token']})[0],400)
        status,result=self.request('entry/apply',{'token':draft['token'],'confirmed':True})
        self.assertEqual(status,200);self.assertEqual(len(result['result']),1)
        self.assertEqual(self.request('entry/apply',{'token':draft['token'],'confirmed':True})[1],result)

    def test_purchase_item_ids_from_browser_are_validated_and_keep_existing_items(self):
        from services.web_workbench_operations import database_fingerprint
        from db.connection import db_read
        previous = self.business.get_record('purchases', 1)
        item = previous['items'][0]
        values = {'item_id': str(item['item_id']), 'product_id': str(item['supplier_offer_id']),
                  'material_name_snapshot': item['material_name_snapshot'],
                  'quantity': str(item['quantity']), 'unit_snapshot': item['unit_snapshot'],
                  'unit_price': self.business.money_form_value(item['material_unit_price_cents']),
                  'tax_percent': str(item['tax_rate_bps'] / 100), 'price_basis': 'exclusive'}
        status, result = self.request('save', {'module': 'purchases', 'id': 1,
            'expected': previous['_version'], 'confirmed': True,
            'record': {**previous, 'notes': '合成原明细备注修改', 'items': [values]}})
        self.assertEqual(status, 200, result)
        current = self.business.get_record('purchases', 1)
        self.assertEqual(current['items'][0]['item_id'], item['item_id'])
        self.assertEqual(current['total_amount_cents'], previous['total_amount_cents'])
        with db_read() as conn:
            before = database_fingerprint(conn)
        for identity in ('1.5', -1, True, 'not-an-id'):
            with self.subTest(identity=identity), self.assertRaisesRegex(ValueError, '编号无效'):
                self.business.prepare_purchase_items([{**values, 'item_id': identity}])
        with db_read() as conn:
            self.assertEqual(database_fingerprint(conn), before)

    def test_all_collections_and_dashboards_have_real_service_results(self):
        for key in self.business.SPECS:
            with self.subTest(module=key):
                status,result=self.request('rows?module='+key)
                self.assertEqual(status,200,result)
                self.assertIsInstance(result['rows'],list)
                for row in result['rows']:
                    self.assertEqual(self.request(f'record?module={key}&id={row["id"]}')[0],200)
        for view in ('home','profit','governance','finance','cost','assistant','collection'):
            self.assertEqual(self.request('dashboard?view='+view)[0],200,view)
        self.assertEqual(self.request('rows?module=unknown')[0],400)
        self.assertEqual(self.request('record?module=unknown&id=1')[0],400)

    def test_form_edits_preserve_existing_amounts_and_version(self):
        for module in ('projects','contracts','settlements','suppliers','customers','offers','purchases','costs','workers','worklogs','construction'):
            with self.subTest(module=module):
                row=self.business.rows(module)[0]
                if module=='suppliers':
                    partner=self.business.partners.get_business_partner(row['id'])
                    self.business.partners.update_business_partner(row['id'],{**partner,'contact_name':'合成供应商联系人','contact_phone':'13800000000','contact_email':'synthetic@example.invalid'})
                original=self.business.get_record(module,row['id'])
                data={f['name']:original.get(f['name'],f['default']) for f in self.business.SPECS[module]['fields']}
                data['notes']='合成网页修改验证'
                if module=='purchases':
                    data['items']=[dict(item,unit_price=str(item['material_unit_price_cents']/100),
                                        tax_percent=str(item['tax_rate_bps']/100)) for item in original['items']]
                self.business.save_record(module,data,row['id'],original['_version'])
                current=self.business.get_record(module,row['id'])
                if module=='suppliers':
                    self.assertEqual(current['contact_phone'],'13800000000')
                    self.assertEqual(current['contact_email'],'synthetic@example.invalid')
                for key in ('amount_minor','tax_inclusive_amount_minor','total_amount_cents','work_amount_cents'):
                    if key in original:
                        self.assertEqual(current[key],original[key],module)
                with self.assertRaises(ValueError):
                    self.business.save_record(module,data,row['id'],'obsolete-version')

    def test_financial_save_requires_confirmation_and_exact_cents(self):
        body={'module':'plans','record':{'direction':'out','title':'合成网页计划','amount':'123456789012.34','due_date':'2026-12-01'}}
        self.assertEqual(self.request('save',body)[0],400)
        status,result=self.request('save',{**body,'confirmed':True})
        self.assertEqual(status,200,result)
        _,record=self.request(f'record?module=plans&id={result["result"]}')
        self.assertEqual(record['amount_minor'],'12345678901234')
        for amount in ('1.234','NaN','-1'):
            self.assertEqual(self.request('save',{**body,'confirmed':True,'record':{**body['record'],'amount':amount}})[0],400)

    def test_host_origin_token_and_read_only_are_enforced(self):
        body={'module':'plans','record':{},'confirmed':True}
        for headers in ({'Origin':'http://external.invalid'}, {'X-Workbench-Token':'wrong'}):
            self.assertEqual(self.request('save',body,headers=headers)[0],403)
        self.assertEqual(self.request('bootstrap',headers={'Host':'external.invalid'})[0],400)
        self.server.read_only=True
        self.assertEqual(self.request('save',body)[0],403)
        self.assertEqual(self.request('bootstrap')[0],200)

    def test_receipts_funds_historical_input_invoice_and_inspection(self):
        receipt=self.business.save_record('receipts',{'project_id':1,'contract_id':1,'settlement_id':1,
            'invoice_id':1,'amount':'100.01','receipt_date':'2026-10-02','payer_name':'合成建设单位'})
        self.assertEqual(self.business.get_record('receipts',receipt)['amount_minor'],10001)
        transaction={'category':'receipt','source_id':receipt,'to_account_id':1,'amount':'100.01',
            'transaction_date':'2026-10-02','request_key':'synthetic-stable-key'}
        first=self.business.save_record('transactions',transaction)
        self.assertEqual(self.business.save_record('transactions',transaction),first)
        self.assertEqual(self.business.funds.list_accounts()[0]['balance_minor'],25010001)
        historical=self.business.save_record('historical',{'customer_id':1,'name':'合成历史项目','opening_amount':'200'})
        self.assertEqual(self.business.get_record('historical',historical)['opening_minor'],20000)
        entity=self.business.entities.list_entities()[0]
        supplier=self.business.partners.get_supplier(2)
        invoice=self.business.save_record('input_invoices',{'entity_id':entity['id'],'supplier_id':2,
            'seller_name':supplier['name'],'buyer_name':entity['name'],'invoice_no':'12345678901234567890',
            'invoice_date':'2026-10-02','amount':'113.01','net_amount':'100.01','tax_amount':'13',
            'tax_rate_label':'13%','deduction_status':'pending'})
        original=self.business.get_record('input_invoices',invoice)
        self.assertEqual(original['net_amount'],'100.01')
        self.business.save_record('input_invoices',{**original,'notes':'合成修改'},invoice,original['_version'])
        self.assertEqual(self.business.get_record('input_invoices',invoice)['gross_minor'],11301)
        self.business.action('inspection',{'id':1,'inspection_status':'已验收','inspection_date':'2026-10-02',
            'inspector':'合成验收人','inspection_notes':'合成验收依据'})
        self.assertEqual(self.business.get_record('construction',1)['inspection_status'],'已验收')

    def test_operation_cancel_confirmation_stale_and_repeat_are_safe(self):
        from services import web_workbench_operations as operations
        body={'name':'fund_reconcile','record':{'account_id':1,'balance_date':'2026-10-02','actual_amount':'249999.99','reason':'合成核对'}}
        _,response=self.request('operation_preview',body)
        token=response['result']['token']
        self.assertEqual(operations.read('reconciliations',{}),[])
        self.assertEqual(response['result']['review']['difference_minor'],'-1')
        self.assertEqual(self.request('operation_apply',{'token':token})[0],400)
        self.request('operation_discard',{'token':token})
        self.assertEqual(self.request('operation_apply',{'token':token,'confirmed':True})[0],400)
        _,response=self.request('operation_preview',body)
        token=response['result']['token']
        self.business.save_record('plans',{'direction':'out','title':'预览之后的合成变更','amount':'1'})
        self.assertEqual(self.request('operation_apply',{'token':token,'confirmed':True})[0],400)
        self.assertEqual(operations.read('reconciliations',{}),[])
        _,response=self.request('operation_preview',body)
        token=response['result']['token']
        first=self.request('operation_apply',{'token':token,'confirmed':True})
        self.assertEqual(first[0],200,first)
        self.assertEqual(self.request('operation_apply',{'token':token,'confirmed':True}),first)
        self.assertEqual(len(operations.read('reconciliations',{})),1)
        self.assertEqual(self.business.funds.list_transactions(),[])
        self.assertEqual(self.business.funds.list_accounts()[0]['balance_minor'],25000000)
        self.assertEqual(self.request('operation_preview',{'name':'receipt_void','record':{}})[0],400)

    def test_operations_preserve_allocation_cents_and_locked_wages(self):
        from datetime import date
        from services import web_workbench_operations as operations
        def execute(name,record):
            status,response=self.request('operation_preview',{'name':name,'record':record})
            self.assertEqual(status,200,response)
            status,response=self.request('operation_apply',{'token':response['result']['token'],'confirmed':True})
            self.assertEqual(status,200,response)
            return response['result']
        other=self.business.save_record('projects',{'name':'合成分摊项目','project_code':'SYNTH-SPLIT'})
        cost=self.business.save_record('costs',{'category':'管理费','amount':'100.01','cost_date':'2026-10-01'})
        execute('cost_allocate',{'id':cost,'method':'equal','project_ids':[1,other],'reason':'合成分摊'})
        lines=operations.read('cost_allocations',{'id':str(cost)})['lines']
        self.assertEqual(sum(row['amount_minor'] for row in lines),10001)
        today=date.today().isoformat()
        locked=self.business.save_record('worklogs',{'worker_id':1,'project_id':1,'construction_site':'合成地点','work_date':today,'work_days':'0.5'})
        unlocked=self.business.save_record('worklogs',{'worker_id':1,'project_id':1,'construction_site':'合成地点','work_date':today,'work_days':'0.5'})
        execute('labor_lock',{'ids':[1,locked],'locked':True,'reason':'合成锁定'})
        values={'worker_id':1,'new_daily_rate':'410.01','effective_from':date.today().isoformat(),'scope_mode':'through_today','reason':'合成调薪'}
        execute('labor_adjust',values)
        self.assertEqual(self.business.get_record('worklogs',1)['amount_minor'],40000)
        self.assertEqual(self.business.get_record('worklogs',1)['rate_locked'],1)
        self.assertEqual(self.business.get_record('worklogs',locked)['amount_minor'],20000)
        self.assertEqual(self.business.get_record('worklogs',unlocked)['amount_minor'],20501)
        self.assertTrue(operations.read('labor_adjust_history',{'id':'1'}))
        execute('labor_lock',{'ids':[1],'locked':False,'reason':'合成解锁'})
        execute('labor_overtime',{'ids':[1],'is_overtime':True,'reason':'合成加班'})
        self.assertTrue(self.business.get_record('worklogs',1)['is_overtime'])

    def test_batch_worklogs_and_collection_followup_are_atomic(self):
        entry={'worker_id':1,'project_id':1,'construction_site':'合成批量地点','work_date':'2026-10-03','work_days':'0.5'}
        before=len(self.business.rows('worklogs'))
        status,preview=self.request('operation_preview',{'name':'labor_batch','record':{'entries':[entry,{**entry,'worker_id':999999}],'reason':'合成错误行'}})
        self.assertEqual(status,400,preview)
        self.assertEqual(len(self.business.rows('worklogs')),before)
        _,preview=self.request('operation_preview',{'name':'labor_batch','record':{'entries':[entry,entry],'reason':'合成正确整批'}})
        status,response=self.request('operation_apply',{'token':preview['result']['token'],'confirmed':True})
        self.assertEqual(status,200,response)
        self.assertEqual(len(self.business.rows('worklogs')),before+2)
        record={'id':1,'owner_name':'合成跟进人','status':'pending','followup_date':'2026-10-03','followup_content':'合成资料待验收','reason':'合成跟进'}
        _,preview=self.request('operation_preview',{'name':'collection_followup','record':record})
        first=self.request('operation_apply',{'token':preview['result']['token'],'confirmed':True})
        self.assertEqual(first[0],200,first)
        self.assertEqual(self.request('operation_apply',{'token':preview['result']['token'],'confirmed':True}),first)
        self.assertEqual(len(self.business.collection.list_followup_logs(1)),1)

    def retire_worker_with_history(self):
        worker=self.business.save_record('workers',{'name':'合成离职工人','daily_rate':'300'})
        log=self.business.save_record('worklogs',{'worker_id':worker,'project_id':1,'construction_site':'合成历史工地','work_date':'2026-09-01','work_days':'0.5'})
        previous=self.business.get_record('workers',worker)
        self.business.save_record('workers',{'status':'离职'},worker,previous['_version'])
        return worker,log

    def test_new_worklog_options_hide_retired_but_archives_and_history_remain(self):
        worker,log=self.retire_worker_with_history()
        status,options=self.request('options')
        self.assertEqual(status,200,options)
        self.assertNotIn(worker,[row['value'] for row in options['active_workers']])
        self.assertIn(worker,[row['value'] for row in options['workers']])
        self.assertIn(worker,[row['id'] for row in self.business.rows('workers')])
        self.assertEqual(self.business.get_record('worklogs',log)['worker_id'],worker)
        matrix=self.business.labor.get_worker_monthly_workdays('2026-09-01','2026-09-30')
        self.assertIn(worker,[row['worker_id'] for row in matrix['workers']])

    def test_retired_worker_cannot_be_new_entry_but_own_history_can_be_edited(self):
        worker,log=self.retire_worker_with_history()
        before=len(self.business.rows('worklogs'))
        entry={'worker_id':worker,'project_id':1,'construction_site':'合成历史工地','work_date':'2026-10-03','work_days':'0.5'}
        status,response=self.request('save',{'module':'worklogs','record':entry,'confirmed':True})
        self.assertEqual(status,400,response)
        self.assertIn('离职',response['error'])
        previous=self.business.get_record('worklogs',log)
        status,response=self.request('save',{'module':'worklogs','id':log,'expected':previous['_version'],'record':{'notes':'合成补充历史说明'},'confirmed':True})
        self.assertEqual(status,200,response)
        self.assertEqual(self.business.get_record('worklogs',log)['notes'],'合成补充历史说明')
        self.assertEqual(len(self.business.rows('worklogs')),before)

    def test_batch_preview_uses_each_worker_rate_and_mixed_days_then_saves_once(self):
        from datetime import date
        today=date.today().isoformat()
        worker=self.business.save_record('workers',{'name':'合成在职工人乙','daily_rate':'200'})
        self.business.labor.apply_rate_adjustment({'worker_id':worker,'new_daily_rate':'300.01','effective_from':today,'scope_mode':'future_only','reason':'合成工资版本'})
        base={'project_id':1,'construction_site':'合成批量出勤','work_date':today}
        entries=[{**base,'worker_id':1,'work_days':'1','is_overtime':False},{**base,'worker_id':worker,'work_days':'0.5','is_overtime':True}]
        before=len(self.business.rows('worklogs'))
        status,preview=self.request('operation_preview',{'name':'labor_batch','record':{'entries':entries,'reason':'合成混合工天'}})
        self.assertEqual(status,200,preview)
        review=preview['result']['review']
        self.assertEqual([int(row['amount_minor']) for row in review['rows']],[40000,15001])
        self.assertEqual(int(review['total_minor']),55001)
        self.assertTrue(all(row['worker_name'] and row['project_name'] for row in review['rows']))
        self.assertEqual(len(self.business.rows('worklogs')),before)
        payload={'token':preview['result']['token'],'confirmed':True}
        self.assertEqual(self.request('operation_apply',{**payload,'confirmed':False})[0],400)
        first=self.request('operation_apply',payload)
        self.assertEqual(first[0],200,first)
        self.assertEqual(self.request('operation_apply',payload),first)
        rows=[row for row in self.business.rows('worklogs') if row['construction_site']=='合成批量出勤']
        self.assertEqual(len(rows),2)
        self.assertEqual(sum(row['work_days'] for row in rows),1.5)
        self.assertEqual(sum(row['work_days'] for row in rows if row['is_overtime']),0.5)
        self.assertEqual(len(self.business.rows('worklogs')),before+2)

    def test_batch_rejects_retired_and_excess_daily_days_before_any_write(self):
        worker,_=self.retire_worker_with_history()
        base={'project_id':1,'construction_site':'合成整批拒绝','work_date':'2026-10-03','work_days':'0.5'}
        before=len(self.business.rows('worklogs'))
        cases=[([{**base,'worker_id':1},{**base,'worker_id':worker}],'离职'),
               ([{**base,'worker_id':1,'work_date':'2026-10-02'}],'不能超过 1')]
        for entries,message in cases:
            with self.subTest(message=message):
                status,response=self.request('operation_preview',{'name':'labor_batch','record':{'entries':entries,'reason':'合成整批拒绝'}})
                self.assertEqual(status,400,response)
                self.assertIn(message,response['error'])
                self.assertEqual(len(self.business.rows('worklogs')),before)

    def test_batch_stale_worker_status_and_cancel_never_save(self):
        worker=self.business.save_record('workers',{'name':'合成临时在职工人','daily_rate':'200'})
        record={'entries':[{'worker_id':worker,'project_id':1,'construction_site':'合成状态变更','work_date':'2026-10-03','work_days':'1'}],'reason':'合成批量'}
        before=len(self.business.rows('worklogs'))
        status,preview=self.request('operation_preview',{'name':'labor_batch','record':record})
        self.assertEqual(status,200,preview)
        token=preview['result']['token']
        self.assertEqual(self.request('operation_discard',{'token':token})[0],200)
        self.assertEqual(self.request('operation_apply',{'token':token,'confirmed':True})[0],400)
        _,preview=self.request('operation_preview',{'name':'labor_batch','record':record})
        previous=self.business.get_record('workers',worker)
        self.business.save_record('workers',{'status':'离职'},worker,previous['_version'])
        status,response=self.request('operation_apply',{'token':preview['result']['token'],'confirmed':True})
        self.assertEqual(status,400,response)
        self.assertIn('变化',response['error'])
        self.assertEqual(len(self.business.rows('worklogs')),before)

    def test_operation_queries_and_full_detail_fields(self):
        from services import web_workbench_operations as operations
        queries={'cost_history':{'id':'1'},'cost_allocations':{'id':'1'},'purchase_history':{'id':'1'},'labor_adjust_history':{'id':'1'},'project_cost_details':{'id':'1'},'sites':{'id':'1'},'fund_sources':{'kind':'purchase'},'reconciliations':{},'entities':{},'entity_records':{'kind':'projects'},'input_links':{'id':'1'},'historical_receipts':{'id':'1'},'collection_history':{'id':'1'},'collection_case':{'id':'1'},'unassigned_labor':{},'unassigned_purchases':{},'fulfillment_gaps':{},'labor_month':{'month':'2026-10'},'worker_month':{'id':'1','month':'2026-10','start':'2026-10-01','end':'2026-10-31'},'customer_business':{'id':'1'},'local_query':{'question':'项目成本是多少','project_id':'1'}}
        for name,parameters in queries.items():
            with self.subTest(query=name):
                from urllib.parse import urlencode
                status,response=self.request('operation_read?'+urlencode({'name':name,**parameters}))
                self.assertEqual(status,200,response)
        settlement=self.business.save_record('settlements',{'project_id':1,'contract_id':1,'amount':'113.01','settlement_date':'2026-10-02','basis':'合成税额测试'})
        invoice=self.business.save_record('invoices',{'project_id':1,'contract_id':1,'settlement_id':settlement,'invoice_no':'SYNTH-TAX-13','amount':'113.01','tax_rate_percent':'13','net_amount':'100.01','tax_amount':'13.00','invoice_date':'2026-10-02'})
        record=self.business.get_record('invoices',invoice)
        self.assertEqual(record['tax_rate_bps'],1300)
        partner=self.business.get_record('customers',1)
        self.business.save_record('customers',{'credit_limit':'123456789012.34','contact_name':'合成联系人','contact_department':'合成部','contact_email':'synthetic@example.invalid'},1,partner['_version'])
        current=self.business.get_record('customers',1)
        self.assertEqual(current['credit_limit_minor'],12345678901234)
        self.assertEqual(current['credit_limit'],'123456789012.34')
        self.assertEqual(current['contact_department'],'合成部')

    def test_entity_record_options_are_named_and_reads_do_not_change_the_ledger(self):
        from db.connection import db_read
        from services.web_workbench_operations import database_fingerprint
        from services import funds_service, operating_entity_service

        receipt = self.business.save_record('receipts', {
            'project_id': 1, 'contract_id': 1, 'amount': '100.00',
            'receipt_date': '2026-10-03', 'payment_method': '银行转账',
        })
        funds_service.record_transaction({
            'request_key': 'synthetic-entity-label-receipt', 'category': 'receipt',
            'source_id': receipt, 'amount': '100.00', 'transaction_date': '2026-10-03',
            'to_account_id': 1,
        })
        with db_read() as conn:
            before = database_fingerprint(conn)
        for kind, (_, label_column, _, _) in operating_entity_service.RECORDS.items():
            with self.subTest(kind=kind):
                status, rows = self.request('operation_read?name=entity_records&kind=' + kind)
                self.assertEqual(status, 200, rows)
                self.assertTrue(rows, kind)
                with db_read() as conn:
                    for row in rows:
                        self.assertIn('date', row)
                        self.assertIn('amount_minor', row)
                        self.assertIn('projects', row)
                        expected = conn.execute(
                            f'SELECT {label_column} FROM {kind} WHERE id=?', (row['id'],)
                        ).fetchone()[0]
                        self.assertEqual(row['label'], str(expected))
                        self.assertTrue(row['label'].strip(), kind)
        with db_read() as conn:
            self.assertEqual(database_fingerprint(conn), before)

    def test_fund_verification_plan_payment_and_input_links_follow_services(self):
        from services import web_workbench_operations as operations
        def execute(name,record):
            status,response=self.request('operation_preview',{'name':name,'record':record})
            self.assertEqual(status,200,response)
            status,response=self.request('operation_apply',{'token':response['result']['token'],'confirmed':True})
            self.assertEqual(status,200,response)
            return response['result']
        order=self.business.procurement.get_purchase_order(1)
        self.business.procurement.update_purchase_order(1,{**order,'purchase_date':'2026-09-30'},order['items'])
        execute('fund_verify',{'kind':'purchase','source_id':1,'settled_amount':'100.01','payee':'合成材料供应商','reason':'合成核实'})
        source=operations.read('fund_sources',{'kind':'purchase'})[0]
        self.assertTrue(source['verified'])
        self.assertEqual(source['opening_settled_minor'],10001)
        plan=self.business.save_record('plans',{'direction':'out','title':'合成采购待付','source_kind':'purchase','source_id':1,'amount':'200.01','due_date':'2026-10-02'})
        transaction={'category':'purchase','source_id':1,'from_account_id':1,'amount':'200.01','transaction_date':'2026-10-02','plan_id':plan,'request_key':'synthetic-payment-key'}
        cash=self.business.save_record('transactions',transaction)
        self.assertEqual(self.business.save_record('transactions',transaction),cash)
        self.assertEqual(operations.read('fund_sources',{'kind':'purchase'})[0]['paid_minor'],20001)
        execute('fund_void_transaction',{'id':cash,'reason':'合成录错更正'})
        self.assertEqual(operations.read('fund_sources',{'kind':'purchase'})[0]['paid_minor'],0)
        entity=self.business.entities.list_entities()[0]
        execute('entity_assign',{'kind':'purchase_orders','ids':[1],'entity_id':entity['id'],'reason':'合成主体确认'})
        supplier=self.business.partners.get_supplier(2)
        invoice=self.business.save_record('input_invoices',{'entity_id':entity['id'],'supplier_id':2,'seller_name':supplier['name'],'buyer_name':entity['name'],'invoice_no':'99887766554433221100','invoice_date':'2026-10-02','amount':'113.01','net_amount':'100.01','tax_amount':'13','tax_rate_label':'13%','deduction_status':'pending'})
        execute('input_links',{'id':invoice,'links':[{'kind':'purchase_orders','record_id':1,'amount':'113.01'}],'reason':'合成票成本核对'})
        links=operations.read('input_links',{'id':str(invoice)})
        self.assertEqual(links[0]['amount_minor'],11301)
        status,response=self.request('operation_preview',{'name':'input_links','record':{'id':invoice,'links':[{'kind':'purchase_orders','record_id':1,'amount':'999999'}],'reason':'合成超额拒绝'}})
        self.assertEqual(status,200,response)
        self.assertEqual(self.request('operation_apply',{'token':response['result']['token'],'confirmed':True})[0],400)
        self.assertEqual(operations.read('input_links',{'id':str(invoice)}),links)

    def test_excel_preview_cancel_atomic_apply_and_duplicate(self):
        from openpyxl import Workbook, load_workbook
        from web_workbench import transfer
        file=Path(self.temp.name)/'synthetic-suppliers.xlsx'
        workbook=Workbook();sheet=workbook.active
        sheet.append(['ID','工厂名称','产品范围','联系人','价格水平','交期','质量情况','是否做出口','备注','默认税率（%）'])
        sheet.append([None,'合成批量供应商','工具','合成联系人','中','一般','良','否','合成测试',13])
        workbook.save(file);workbook.close()
        upload={'module':'suppliers','name':file.name,'content':base64.b64encode(file.read_bytes()).decode()}
        count=len(self.business.rows('suppliers'))
        status,response=self.request('excel_preview',upload)
        self.assertEqual(status,200,response)
        self.assertEqual(len(self.business.rows('suppliers')),count)
        token=response['result']['token'];self.request('excel_discard',{'token':token})
        self.assertEqual(self.request('excel_apply',{'token':token,'confirmed':True})[0],400)
        _,response=self.request('excel_preview',upload)
        token=response['result']['token']
        self.assertEqual(self.request('excel_apply',{'token':token})[0],400)
        first=self.request('excel_apply',{'token':token,'confirmed':True})
        self.assertEqual(first[0],200,first)
        self.assertEqual(self.request('excel_apply',{'token':token,'confirmed':True}),first)
        self.assertEqual(len(self.business.rows('suppliers')),count+1)
        self.assertEqual(self.request('excel_preview',upload)[0],400)
        broken=load_workbook(file);broken.active.append([None,'另一个合成供应商','工具','人','中','一般','良','否','合成测试',-13])
        broken.active.cell(2,2).value='应整批回滚的合成供应商';broken.save(file);broken.close()
        upload['content']=base64.b64encode(file.read_bytes()).decode()
        self.assertEqual(self.request('excel_preview',upload)[0],400)
        self.assertFalse(any(row['name']=='应整批回滚的合成供应商' for row in self.business.rows('suppliers')))
        target=Path(self.temp.name)/'synthetic-export.xlsx'
        transfer.run(self.path,'export','operating',target)
        with __import__('contextlib').closing(load_workbook(target)) as exported:
            self.assertIn('完整字段与精确分值',exported.sheetnames)
            fields=exported['完整字段与精确分值']
            self.assertTrue(any(row[2]=='amount_minor' and isinstance(row[3],str) for row in fields.iter_rows(min_row=2,values_only=True)))

    def test_excel_purchase_multiline_freight_formulas_and_offer_precision(self):
        from contextlib import closing
        from openpyxl import Workbook,load_workbook
        from web_workbench import transfer
        file=Path(self.temp.name)/'synthetic-purchases.xlsx'
        transfer.run(self.path,'export','purchases',file)
        with closing(load_workbook(file)) as workbook:
            sheet=workbook.active;row=list(next(sheet.iter_rows(min_row=2,max_row=2,values_only=True)))
            row[0]=None;row[1]='SYNTH-EXCEL-MULTILINE';row[18]='100.01'
            for index,value in enumerate(row,start=1):sheet.cell(2,index).value=value
            second=list(row);second[18]='0';sheet.append(second)
            workbook.save(file)
        payload={'module':'purchases','name':file.name,'content':base64.b64encode(file.read_bytes()).decode()}
        status,response=self.request('excel_preview',payload)
        self.assertEqual(status,200,response)
        self.assertEqual(response['result']['review']['row_count'],2)
        self.assertEqual(response['result']['review']['record_count'],1)
        token=response['result']['token']
        status,response=self.request('excel_apply',{'token':token,'confirmed':True})
        self.assertEqual(status,200,response)
        imported=self.business.get_record('purchases',response['result']['rows'][0]['id'])
        self.assertEqual(len(imported['items']),2)
        self.assertEqual(imported['freight_amount_cents'],10001)
        self.assertEqual(imported['total_amount_cents'],sum(row['line_amount_cents'] for row in imported['items'])+10001)
        self.assertEqual(self.request('excel_preview',payload)[0],400)
        with closing(load_workbook(file)) as workbook:
            workbook.active.cell(2,2).value=None;workbook.active.cell(3,2).value=None;workbook.save(file)
        missing_number={**payload,'content':base64.b64encode(file.read_bytes()).decode()}
        status,response=self.request('excel_preview',missing_number)
        self.assertEqual(status,400,response)
        self.assertIn('请补采购单号',response['error'])
        with closing(load_workbook(file)) as workbook:
            workbook.active.cell(2,13).value='=1+1';workbook.save(file)
        payload['content']=base64.b64encode(file.read_bytes()).decode()
        self.assertEqual(self.request('excel_preview',payload)[0],400)
        workbook=Workbook();sheet=workbook.active
        sheet.append(['ID','供应商ID','供应商名称','产品名称','规格','单位','材料单价（未税）','税率（%）','含税单价','备注','报价口径'])
        sheet.append([None,2,'合成供应商','合成超大额材料','合成规格','件','90071992547409.93',0,None,'合成精度测试','exclusive'])
        workbook.save(file);workbook.close()
        payload={'module':'offers','name':file.name,'content':base64.b64encode(file.read_bytes()).decode()}
        status,response=self.request('excel_preview',payload)
        self.assertEqual(status,200,response)
        self.assertEqual(response['result']['review']['rows'][0]['quoted_price_minor'],'9007199254740993')

    def test_purchase_weight_total_inclusive_price_and_pdf_unknowns(self):
        from pypdf import PdfWriter
        from services import web_workbench_operations as operations
        for mode,expected in [('weight',225001),('total',123456)]:
            item={'product_id':1,'material_name_snapshot':'合成计价材料','unit_snapshot':'吨','quantity':'1','unit_price':'113.01','price_basis':'inclusive','tax_percent':'13','settlement_mode':mode,'net_weight':'0.5','weight_unit':'吨','weight_unit_price':'4500.01','settlement_total':'1234.56'}
            order=self.business.save_record('purchases',{'project_id':1,'supplier_id':2,'purchase_date':'2026-10-03','freight_amount':'0.01','items':[item]})
            original=self.business.get_record('purchases',order)
            self.assertEqual(original['total_amount_cents'],expected+1)
            data={**original,'items':[{**row,'unit_price':self.business.money_form_value(row['tax_inclusive_unit_price_cents']),'tax_percent':'13','settlement_total':self.business.money_form_value(row['settlement_total_cents'] or 0)} for row in original['items']]}
            self.business.save_record('purchases',data,order,original['_version'])
            self.assertEqual(self.business.get_record('purchases',order)['total_amount_cents'],expected+1)
            self.assertTrue(operations.read('purchase_history',{'id':str(order)}))
        offer=self.business.get_record('offers',1)
        self.business.save_record('offers',{'price':'4501.01'},1,offer['_version'])
        self.assertTrue(operations.read('offer_history',{'id':'1'}))
        pdf=Path(self.temp.name)/'synthetic-blank.pdf'
        writer=PdfWriter();writer.add_blank_page(100,100);writer.write(pdf);writer.close()
        upload={'name':pdf.name,'content':base64.b64encode(pdf.read_bytes()).decode()}
        status,response=self.request('input_pdf_preview',upload)
        self.assertEqual(status,400,response)
        self.assertIn('扫描PDF',response['error'])
        self.assertEqual(self.business.rows('input_invoices'),[])
        self.assertEqual(self.server.uploads,{})

    def cashflow_attachment(self):
        from services import attachment_service as attachments
        path = Path(self.temp.name) / (self.id().split('.')[-1] + '.txt')
        path.write_text('收款 合成尾款；金额：100元；日期：2026-12-01；条件：合成验收', encoding='utf-8')
        with patch.object(attachments, '_storage_root', return_value=Path(self.temp.name) / 'source-files'):
            identity = attachments.add_attachment('contract', 1, path, category='合同附件')
        return attachments.list_attachments('contract', 1)[0], identity

    def test_attachment_preview_cancel_confirm_and_scope_bound_retry(self):
        from services import document_cashflow_service as cashflow
        attachment, identity = self.cashflow_attachment()
        scope = {'contract_id': 1, 'project_id': 1}
        _, rows = self.request('cashflow/attachments?contract_id=1&project_id=1')
        self.assertEqual(rows[0]['id'], identity)
        self.assertTrue(rows[0]['file_exists'])
        self.assertFalse(rows[0]['already_linked'])
        count = len(cashflow.list_documents(**scope))
        status, response = self.request('cashflow/attachment_preview', {**scope, 'attachment_id': identity})
        self.assertEqual(status, 200, response)
        token = response['result']['token']
        self.assertEqual(len(cashflow.list_documents(**scope)), count)
        self.request('cashflow/discard', {'token': token})
        self.assertTrue(Path(attachment['absolute_path']).is_file())
        self.assertEqual(self.request('cashflow/link_attachment', {**scope, 'token': token, 'kind': '合同正文', 'acknowledged': True})[0], 400)
        _, response = self.request('cashflow/attachment_preview', {**scope, 'attachment_id': identity})
        data = {**scope, 'token': response['result']['token'], 'kind': '合同正文'}
        self.assertEqual(self.request('cashflow/link_attachment', data)[0], 400)
        data['acknowledged'] = True
        self.assertEqual(self.request('cashflow/link_attachment', {**data, 'project_id': None})[0], 400)
        first = self.request('cashflow/link_attachment', data)
        self.assertEqual(first[0], 200, first)
        self.assertEqual(self.request('cashflow/link_attachment', data), first)
        self.assertEqual(len(cashflow.list_documents(**scope)), count + 1)
        self.assertEqual(self.request('cashflow/link_attachment', {**data, 'kind': '预算'})[0], 400)
        self.assertTrue(Path(attachment['absolute_path']).is_file())
        _, rows = self.request('cashflow/attachments?contract_id=1&project_id=1')
        self.assertTrue(rows[0]['already_linked'])
        self.assertEqual(self.request('cashflow/attachments?contract_id=1')[1][0]['already_linked'], False)
        self.assertEqual(len(self.business.funds.list_transactions()), 0)

    def test_model_source_preview_is_scope_bound_trusted_and_not_bookkeeping(self):
        from unittest.mock import Mock
        from tests.test_document_understanding import RAW_TEXT, interpretation
        from services import document_cashflow_service as cashflow
        scope = {'contract_id': 1, 'project_id': 1}
        upload = {**scope, 'name': '合成模型合同.txt', 'content': base64.b64encode(RAW_TEXT.encode()).decode(),
                  'use_model': True, 'kind': '合同正文'}
        before_documents = len(cashflow.list_documents(**scope))
        before_transactions = len(self.business.funds.list_transactions())
        before_plans = len(self.business.funds.list_plans())
        self.assertEqual(self.request('cashflow/preview', {**upload, 'use_model': 'true'})[0], 400)
        self.assertEqual(self.request('cashflow/preview', {**upload, 'contract_id': None, 'project_id': None})[0], 400)
        client = Mock(model='synthetic-vision', document_understanding=Mock(return_value=interpretation()))
        with patch('ai_engine.make_ai_client', return_value=client):
            status, response = self.request('cashflow/preview', upload)
        self.assertEqual(status, 200, response)
        client.document_understanding.assert_called_once()
        preview = response['result']
        self.assertEqual(preview['understanding']['nodes'][1]['amount_minor'], '4000000')
        self.assertEqual(len(cashflow.list_documents(**scope)), before_documents)
        data = {**scope, 'token': preview['token'], 'kind': '合同正文'}
        self.assertEqual(self.request('cashflow/save_document', data)[0], 400)
        # Posted fields cannot replace the server's original interpretation.
        data.update(acknowledged=True, fragments=[{'locator': '伪造', 'text': '付款；金额：999元'}])
        with patch('services.attachment_service._storage_root', return_value=Path(self.temp.name) / 'model-files'):
            status, response = self.request('cashflow/save_document', data)
        self.assertEqual(status, 200, response)
        status, response = self.request('cashflow/analyze', scope)
        self.assertEqual(status, 200, response)
        _, analysis = self.request('cashflow/analysis?id=' + str(response['result']))
        own_rows = [row for row in analysis['payload']['rows']
                    if any(source['name'] == upload['name'] for source in row['sources'])]
        self.assertEqual(len(own_rows), 3)
        self.assertEqual(len(analysis['payload']['interpretations']), 1)
        self.assertEqual(len(self.business.funds.list_transactions()), before_transactions)
        self.assertEqual(len(self.business.funds.list_plans()), before_plans)

    def test_manual_source_cannot_forge_semantic_packet(self):
        from tests.test_document_understanding import RAW_TEXT, packet
        from services import document_cashflow_service as cashflow
        scope = {'contract_id': 1, 'project_id': 1}
        status, response = self.request('cashflow/save_document', {**scope, 'name': '合成人工资料', 'kind': '人工补充',
            'fragments': [{'locator': '第1行', 'text': RAW_TEXT, 'understanding': packet()}]})
        self.assertEqual(status, 200, response)
        saved = cashflow.list_documents(**scope)[0]
        self.assertNotIn('understanding', json.loads(saved['fragments_json'])[0])
        for malformed in (None, {}, ['invalid-fragment']):
            with self.subTest(malformed=malformed):
                status, _ = self.request('cashflow/save_document', {**scope,
                    'name': '合成无效资料', 'kind': '人工补充', 'fragments': malformed})
                self.assertEqual(status, 400)

    def test_model_read_does_not_hold_write_lock_and_rejects_changed_context(self):
        from unittest.mock import Mock
        from tests.test_document_understanding import RAW_TEXT, interpretation
        from services.document_understanding_service import understand_document
        from db.connection import db_transaction
        entered, release = threading.Event(), threading.Event()
        client = Mock(model='synthetic-vision', document_understanding=Mock(return_value=interpretation()))
        def slow_read(path, **kwargs):
            entered.set()
            if not release.wait(5):
                raise ValueError('合成超时')
            return understand_document(path, client=client, **kwargs)
        responses = []
        data = {'contract_id': 1, 'project_id': 1, 'use_model': True, 'name': '合成慢请求.txt',
                'content': base64.b64encode(RAW_TEXT.encode()).decode()}
        with patch('web_workbench.cashflow_sources.understand_document', side_effect=slow_read):
            worker = threading.Thread(target=lambda: responses.append(self.request('cashflow/preview', data)))
            worker.start()
            try:
                self.assertTrue(entered.wait(2))
                self.assertTrue(self.server.write_lock.acquire(timeout=.5))
                self.server.write_lock.release()
                self.assertEqual(self.request('cashflow/attachments?contract_id=1&project_id=1')[0], 200)
                with db_transaction() as conn:
                    conn.execute("UPDATE projects SET name='合成项目已更名' WHERE id=1")
            finally:
                release.set()
                worker.join(5)
        self.assertEqual(responses[0][0], 400, responses)
        self.assertIn('在理解期间已变化', responses[0][1]['error'])
        self.assertFalse(any(ticket.get('cashflow') for ticket in self.server.uploads.values()))

    def test_attachment_preview_rejects_changed_missing_expired_or_retired_source(self):
        from services import attachment_service as attachments
        attachment, identity = self.cashflow_attachment()
        scope = {'contract_id': 1, 'project_id': 1}
        def preview():
            status, response = self.request('cashflow/attachment_preview', {**scope, 'attachment_id': identity})
            self.assertEqual(status, 200, response)
            return {**scope, 'token': response['result']['token'], 'kind': '合同正文', 'acknowledged': True}
        data = preview()
        self.server.uploads[data['token']]['expires'] = 0
        self.assertEqual(self.request('cashflow/link_attachment', data)[0], 400)
        data = preview()
        file = Path(attachment['absolute_path'])
        file.write_text('合成修改', encoding='utf-8')
        self.assertEqual(self.request('cashflow/link_attachment', data)[0], 400)
        data = preview()
        file.unlink()
        self.assertEqual(self.request('cashflow/link_attachment', data)[0], 400)
        file.write_text('合成修改', encoding='utf-8')
        data = preview()
        attachments.void_attachments([identity])
        self.assertEqual(self.request('cashflow/link_attachment', data)[0], 400)
        self.assertTrue(file.is_file())

    def test_attachment_ocr_preview_does_not_block_other_pages_or_write_in_readonly(self):
        from services.document_cashflow_extract import read_document
        attachment, identity = self.cashflow_attachment()
        scope = {'contract_id': 1, 'project_id': 1}
        original = read_document(attachment['absolute_path'])
        entered, release = threading.Event(), threading.Event()
        def slow_read(path):
            entered.set()
            if not release.wait(5):
                raise ValueError('合成超时')
            return original
        self.server.read_only = True
        responses = []
        with patch('web_workbench.cashflow_sources.read_document', side_effect=slow_read):
            worker = threading.Thread(target=lambda: responses.append(self.request('cashflow/attachment_preview',
                {**scope, 'attachment_id': identity})))
            worker.start()
            try:
                self.assertTrue(entered.wait(2))
                self.assertTrue(self.server.write_lock.acquire(timeout=.5))
                self.server.write_lock.release()
                self.assertEqual(self.request('cashflow/attachments?contract_id=1&project_id=1')[0], 200)
            finally:
                release.set()
                worker.join(5)
        self.assertEqual(responses[0][0], 200, responses)
        token = responses[0][1]['result']['token']
        self.assertEqual(self.request('cashflow/link_attachment', {**scope, 'token': token, 'kind': '合同正文', 'acknowledged': True})[0], 403)
        self.assertEqual(self.request('cashflow/discard', {'token': token})[0], 200)
        self.assertTrue(Path(attachment['absolute_path']).is_file())

    def test_upload_cancel_duplicate_conflict_stale_and_confirm(self):
        scope={'project_id':1,'contract_id':1}
        original_count=len(self.business.funds.list_plans())
        content='付款 合成网页节点；金额：100.01元；日期：2026-11-20；条件：合成验收'.encode()
        upload={'name':'合成网页.txt','content':base64.b64encode(content).decode()}
        status,preview=self.request('cashflow/preview',upload)
        self.assertEqual(status,200,preview)
        token=preview['result']['token']
        self.assertEqual(self.request('cashflow/discard',{'token':token})[0],200)
        self.assertEqual(self.request('cashflow/save_document',{**scope,'token':token,'kind':'采购付款安排'})[0],400)
        with patch('services.attachment_service._storage_root',return_value=Path(self.temp.name)/'attachments'):
            for expected in (False,True):
                _,preview=self.request('cashflow/preview',upload)
                status,result=self.request('cashflow/save_document',{**scope,'token':preview['result']['token'],'kind':'采购付款安排'})
                self.assertEqual(status,200,result)
                self.assertEqual(result['result']['duplicate'],expected)
        _,analysis=self.request('cashflow/analyze',scope)
        analysis_id=analysis['result']
        _,result=self.request('cashflow/analysis?id='+str(analysis_id))
        node=next(r for r in result['payload']['rows'] if r['title']=='合成网页节点')
        confirm={'id':analysis_id,'keys':[node['key']],'reason':'合成浏览器端核对'}
        self.assertEqual(self.request('cashflow/confirm',confirm)[0],400)
        confirm['acknowledged']=True
        first=self.request('cashflow/confirm',confirm)
        self.assertEqual(first[0],200,first)
        self.assertEqual(self.request('cashflow/confirm',confirm),first)
        self.assertEqual(len(self.business.funds.list_plans()),original_count+1)
        self.assertEqual(len(self.business.funds.list_transactions()),0)
        _,new_analysis=self.request('cashflow/analyze',scope)
        self.request('cashflow/save_document',{**scope,'name':'合成网页修订','kind':'预算','fragments':[
            {'locator':'第2条','text':'付款 合成网页节点；金额：110元；日期：2026-11-20；条件：合成验收'}]})
        _,old=self.request('cashflow/analysis?id='+str(new_analysis['result']))
        self.assertTrue(old['stale_reason'])
        self.assertEqual(self.request('cashflow/confirm',{**confirm,'id':new_analysis['result']})[0],400)
        _,new=self.request('cashflow/analyze',scope)
        _,new=self.request('cashflow/analysis?id='+str(new['result']))
        self.assertTrue(next(r for r in new['payload']['rows'] if r['title']=='合成网页节点')['conflict'])
        self.assertEqual(self.request('cashflow/cancel',{'id':new['id']})[0],200)

if __name__=='__main__':
    unittest.main()
