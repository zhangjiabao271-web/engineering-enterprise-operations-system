"""Excel adapters preserve atomicity, exact cents and existing templates, without widgets."""
from contextlib import closing
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from db import connection, migration_runner
from db.schema import init_db
from services import excel_transfer_service as excel, project_service
from services import web_workbench_service as business
from services.web_workbench_operations import database_fingerprint
from web_workbench import transfer
from web_workbench.synthetic import seed


class ExcelTransferServiceTests(unittest.TestCase):
    def setUp(self):
        directory = tempfile.TemporaryDirectory(prefix='headless_excel_')
        self.addCleanup(directory.cleanup)
        self.folder = Path(directory.name)
        self.database = self.folder / 'synthetic.db'
        for module in (connection, migration_runner):
            selected = patch.object(module, 'DB_PATH', self.database)
            selected.start()
            self.addCleanup(selected.stop)
        init_db()
        seed()

    def fingerprint(self):
        with connection.db_read() as conn:
            return database_fingerprint(conn)

    def file(self, headers, rows):
        path = self.folder / 'import.xlsx'
        workbook = Workbook()
        workbook.active.append(headers)
        for row in rows:
            workbook.active.append(row)
        workbook.save(path)
        workbook.close()
        return path

    def suppliers(self, rows):
        return self.file(['ID','工厂名称','产品范围','联系人','价格水平','交期',
                          '质量情况','是否做出口','备注','默认税率（%）'], rows)

    def test_preview_is_read_only_apply_is_atomic_and_duplicate_is_rejected(self):
        path = self.suppliers([[None,'合成导入商户','工具','','中','一般','良','否','',13]])
        before = self.fingerprint()
        review = transfer.worker('preview', 'suppliers', path, '')
        self.assertEqual(before, self.fingerprint())
        self.assertEqual(review['row_count'], 1)
        saved = transfer.worker('apply', 'suppliers', path, review['fingerprint'])
        self.assertEqual(saved['record_count'], 1)
        after = self.fingerprint()
        with self.assertRaisesRegex(ValueError, '已存在'):
            transfer.worker('preview', 'suppliers', path, '')
        self.assertEqual(after, self.fingerprint())

    def test_invalid_late_row_rolls_back_prior_rows(self):
        path = self.suppliers([[None,'合成必须回滚','','','','','','','',13],
                               [None,'合成无效税率','','','','','','','',-1]])
        before = self.fingerprint()
        with self.assertRaises(ValueError):
            transfer.worker('preview', 'suppliers', path, '')
        with self.assertRaises(ValueError):
            transfer.worker('apply', 'suppliers', path, before)
        self.assertEqual(before, self.fingerprint())

    def test_missing_name_on_nonempty_row_is_not_silently_skipped(self):
        path = self.suppliers([[None,None,'非空行']])
        with self.assertRaisesRegex(ValueError, '第 2 行'):
            transfer.worker('preview', 'suppliers', path, '')

    def test_stale_review_refuses_commit(self):
        path = self.suppliers([[None,'合成迟到导入','','','','','','','',0]])
        review = transfer.worker('preview', 'suppliers', path, '')
        project_service.create_project({'name':'合成预览后新业务'})
        before = self.fingerprint()
        with self.assertRaisesRegex(ValueError, '已变化'):
            transfer.worker('apply', 'suppliers', path, review['fingerprint'])
        self.assertEqual(before, self.fingerprint())

    def test_shared_service_transaction_does_not_commit_project_creation(self):
        before = self.fingerprint()
        with self.assertRaisesRegex(ValueError, 'rollback'):
            with connection.db_transaction(immediate=True) as conn, transfer.bind_business_connections(conn):
                project_service.create_project({'name':'合成共享事务回滚'})
                raise ValueError('rollback')
        self.assertEqual(before, self.fingerprint())

    def test_supplier_default_rate_and_exact_large_offer_price(self):
        supplier = business.save_record('suppliers', {'name':'合成默认税率厂','default_tax_rate_percent':'13'})
        path = self.file(['ID','供应商ID','供应商名称','产品名称','规格','单位','材料单价（未税）','备注'],
                         [[None,supplier,'合成默认税率厂','合成精确材料','A','件','90071992547409.93','']])
        review = transfer.worker('preview', 'offers', path, '')
        self.assertEqual(review['rows'][0]['quoted_price_minor'], 9007199254740993)
        self.assertEqual(review['rows'][0]['tax_rate_bps'], 1300)

    def test_exports_do_not_turn_business_names_into_formulas(self):
        business.save_record('suppliers', {'name':'=SUM(1,2)'})
        path = self.folder / 'safe-export.xlsx'
        transfer.worker('export', 'suppliers', path, '')
        with closing(load_workbook(path)) as workbook:
            cells = [cell for row in workbook['供应商'] for cell in row if cell.value == '=SUM(1,2)']
            self.assertEqual(len(cells), 1)
            self.assertEqual(cells[0].data_type, 's')

    def test_decimal_guards_keep_nonfinite_values_out(self):
        for value in ('NaN', 'Infinity', '-1'):
            with self.subTest(value=value), self.assertRaises(ValueError):
                excel.import_decimal(value, '金额')
        self.assertEqual(excel.import_minor(excel.import_decimal('1.005', '金额')), 101)


if __name__ == '__main__':
    unittest.main()
