"""Exercise weight/total UI, attachments and Excel on an isolated snapshot."""
import os
import sys
import tempfile
import time
from pathlib import Path
from unittest.mock import patch


def descendants(widget):
    for child in widget.winfo_children():
        yield child
        yield from descendants(child)


def field(dialog, text):
    label = next(w for w in descendants(dialog) if hasattr(w, "cget") and w.winfo_class() == "TLabel" and w.cget("text") == text)
    position = label.grid_info()
    return label.master.grid_slaves(row=position["row"], column=position["column"] + 1)[0]


def main():
    project_root = Path(__file__).resolve().parents[1]
    sys.path.insert(0, str(project_root))
    with tempfile.TemporaryDirectory(prefix="weight_ui_") as folder:
        os.environ["SUPPLY_CHAIN_DB_PATH"] = str(Path(folder)/"test.db")
        os.environ["SUPPLY_CHAIN_ATTACHMENTS_PATH"] = str(Path(folder)/"attachments")
        from db.backup import backup_database
        from db.migration_runner import run_migrations
        backup_database(project_root/"supplier_data.db", os.environ["SUPPLY_CHAIN_DB_PATH"])
        run_migrations()
        from services import master_data_service as master, project_service, procurement_service as purchase, attachment_service
        supplier = master.create_supplier({"name": "过磅界面验收供应商"})
        project = project_service.create_project({"name": "钢平台过磅验收项目"})
        header = {"purchase_type": "正式采购", "supplier_id": supplier, "project_id": project, "merchant_name_snapshot": "过磅界面验收供应商", "purchase_date": "2026-09-09", "freight_amount_cents": 10000}
        item = {"material_name_snapshot": "混合规格钢材", "specification_snapshot": "见规格清单", "settlement_mode": "weight", "net_weight": "5.26", "weight_unit": "吨", "weight_unit_price": "3800", "price_basis": "inclusive", "tax_rate_bps": 1300}
        order = purchase.add_purchase_order(header, item)
        import ttkbootstrap as ttk
        from pages.purchase_management_page import PurchaseManagementPage
        from pages.import_export_page import ImportExportPage
        from ui.theme import configure_design_system
        root = ttk.Window(themename="flatly")
        root.geometry("1366x768")
        configure_design_system(root)
        errors = []
        root.report_callback_exception = lambda *args: errors.append(str(args))
        try:
            page = PurchaseManagementPage(root)
            page.open_purchase_dialog("正式采购", order)
            root.update()
            dialog = next(w for w in descendants(root) if isinstance(w, ttk.Toplevel))
            entries = [w for w in descendants(dialog) if isinstance(w, ttk.Entry)]
            assert "20088.00" in [w.get() for w in entries]
            assert any(w.get() == "5.26" and w.winfo_ismapped() for w in entries)
            dialog.attributes("-topmost", True)
            dialog.lift()
            root.update()
            time.sleep(.3)
            from PIL import ImageGrab
            output = project_root/"qa/review_20260906/weight_purchase.png"
            ImageGrab.grab(bbox=(dialog.winfo_rootx(), dialog.winfo_rooty(), dialog.winfo_rootx()+dialog.winfo_width(), dialog.winfo_rooty()+dialog.winfo_height())).save(output)
            save = next(w for w in descendants(dialog) if isinstance(w, ttk.Button) and w.cget("text") == "保存修改")
            with patch("pages.purchase_management_page.messagebox.showwarning") as warning:
                save.invoke()
                root.update()
                assert not warning.called, warning.call_args
            assert purchase.get_purchase_order(order)["net_weight"] == "5.26"
            transfer = ImportExportPage.__new__(ImportExportPage)
            transfer.parent = root
            path = str(Path(folder)/"weight.xlsx")
            row = purchase.get_purchase_order(order)
            with patch("pages.import_export_page.filedialog.asksaveasfilename", return_value=path), patch("pages.import_export_page.messagebox.showinfo"), patch.object(purchase, "list_purchase_orders", return_value=[row]):
                transfer.export_purchases()
            from openpyxl import load_workbook
            book = load_workbook(path)
            book.active.cell(2,2,"WEIGHT-ROUNDTRIP")
            book.save(path)
            book.close()
            with patch("pages.import_export_page.filedialog.askopenfilename", return_value=path), patch("pages.import_export_page.messagebox.showinfo") as info:
                transfer.import_purchases()
                assert info.called
            imported = purchase.list_purchase_orders(keyword="WEIGHT-ROUNDTRIP")[0]
            assert imported["net_weight"] == "5.26" and imported["project_cost_cents"] == 2008800
            page.open_purchase_dialog("正式采购", order)
            root.update()
            dialog = next(w for w in descendants(root) if isinstance(w, ttk.Toplevel))
            mode = next(w for w in descendants(dialog) if isinstance(w, ttk.Combobox) and "按结算总额" in w["values"])
            mode.set("按结算总额")
            root.update()
            label = next(w for w in descendants(dialog) if isinstance(w, ttk.Label) and w.cget("text") == "整批结算金额（不含运费）*")
            entry = field(dialog, "整批结算金额（不含运费）*")
            assert entry.winfo_ismapped()
            entry.insert(0, "8888.88")
            save = next(w for w in descendants(dialog) if isinstance(w, ttk.Button) and w.cget("text") == "保存修改")
            with patch("pages.purchase_management_page.messagebox.showwarning") as warning:
                save.invoke()
                root.update()
                assert not warning.called, warning.call_args
            assert purchase.get_purchase_order(order)["project_cost_cents"] == 898888
            page.formal_tree.selection_set(str(order))
            page.open_purchase_attachments(page.formal_tree)
            root.update()
            attachment_dialog = next(w for w in descendants(root) if isinstance(w, ttk.Toplevel))
            source = Path(folder)/"weigh.txt"
            source.write_text("weigh test", encoding="utf-8")
            add = next(w for w in descendants(attachment_dialog) if isinstance(w, ttk.Button) and w.cget("text") == "添加文件")
            with patch("ui.attachments.filedialog.askopenfilename", return_value=str(source)):
                add.invoke()
                root.update()
            assert len(attachment_service.list_attachments("purchase", order)) == 1
            attachment_dialog.destroy()
            page.open_purchase_dialog("正式采购")
            root.update()
            dialog = next(w for w in descendants(root) if isinstance(w, ttk.Toplevel))
            field(dialog, "计价方式").set("按过磅重量")
            supplier_combo = field(dialog, "供应商 *")
            supplier_combo.set("过磅界面验收供应商")
            supplier_combo.event_generate("<<ComboboxSelected>>")
            project_combo = field(dialog, "所属项目")
            project_combo.set(next(value for value in project_combo["values"] if "钢平台过磅验收项目" in value))
            field(dialog, "材料 *").set("新录整车钢材")
            field(dialog, "实际净重 *").insert(0, "5.26")
            field(dialog, "重量单价（含税，元/吨）*").insert(0, "3800")
            tax = field(dialog, "税率（%）*")
            tax.delete(0, "end")
            tax.insert(0, "13")
            root.update()
            save = next(w for w in descendants(dialog) if isinstance(w, ttk.Button) and w.cget("text") == "保存并关闭")
            with patch("pages.purchase_management_page.messagebox.showwarning") as warning:
                save.invoke()
                root.update()
                assert not warning.called, warning.call_args
            fresh = purchase.list_purchase_orders(keyword="新录整车钢材")
            assert len(fresh) == 1 and fresh[0]["project_cost_cents"] == 1998800 and fresh[0]["product_id"] is None
            assert not errors, errors
            print("PASS: weight preview/save, mixed material without offer, Excel roundtrip, switch to total, purchase attachment upload, new formal weight purchase")
        finally:
            root.destroy()


if __name__ == "__main__":
    main()
