"""Desktop acceptance of inclusive price offers and procurement, on a snapshot."""
import os
import sys
import tempfile
from pathlib import Path
from unittest.mock import patch


def descendants(widget):
    for child in widget.winfo_children():
        yield child
        yield from descendants(child)


def main():
    root_dir = Path(__file__).resolve().parents[1]
    sys.path.insert(0, str(root_dir))
    with tempfile.TemporaryDirectory(prefix="inclusive_ui_") as folder:
        os.environ["SUPPLY_CHAIN_DB_PATH"] = str(Path(folder) / "test.db")
        sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
        from db.backup import backup_database
        from db.migration_runner import run_migrations
        backup_database(root_dir / "supplier_data.db", os.environ["SUPPLY_CHAIN_DB_PATH"])
        run_migrations()
        import ttkbootstrap as ttk
        from services import procurement_service as purchase, master_data_service as master, project_service
        from pages.product_page import ProductPage
        from pages.purchase_management_page import PurchaseManagementPage
        from ui.theme import configure_design_system

        supplier = master.create_supplier({"name": "含税界面测试供应商"})
        offer = master.create_supplier_offer({"supplier_id": supplier, "name": "界面验收螺栓", "specification": "20*55", "unit": "套", "price": "1.89", "price_basis": "inclusive", "tax_rate_percent": 13})
        project = project_service.create_project({"name": "含税界面验收项目"})
        header = {"purchase_type": "正式采购", "project_id": project, "supplier_id": supplier, "merchant_name_snapshot": "含税界面测试供应商", "purchase_date": "2026-09-06"}
        item = {"product_id": offer, "material_name_snapshot": "界面验收螺栓", "specification_snapshot": "20*55", "quantity": 996, "tax_rate_bps": 1300, "price_basis": "inclusive", "tax_inclusive_unit_price_cents": 189}
        order = purchase.add_purchase_order(header, item)
        root = ttk.Window(themename="flatly")
        root.geometry("1366x768")
        configure_design_system(root)
        errors = []
        root.report_callback_exception = lambda *args: errors.append(str(args))
        try:
            host = ttk.Frame(root)
            host.pack(fill="both", expand=True)
            products = ProductPage(host)
            products.set_form_data(master.get_supplier_offer(offer))
            root.update()
            assert products.price_basis_var.get() == "含税价"
            assert products.price_entry.get() == "1.89"
            assert products.tax_inclusive_price_var.get() == "1.89"
            master.update_supplier_offer(offer, products.get_form_data())
            assert master.get_supplier_offer(offer)["quoted_price"] == 1.89
            host.destroy()
            host = ttk.Frame(root)
            host.pack(fill="both", expand=True)
            page = PurchaseManagementPage(host)
            page.open_purchase_dialog("正式采购", order)
            root.update()
            dialog = next(w for w in descendants(root) if isinstance(w, ttk.Toplevel))
            dialog.attributes("-topmost", True)
            dialog.lift()
            import time
            root.update()
            time.sleep(0.4)
            root.update()
            entries = [w.get() for w in descendants(dialog) if isinstance(w, ttk.Entry)]
            assert "1882.44" in entries and "1.89" in entries, entries
            assert any(isinstance(w, ttk.Combobox) and w.get() == "含税价" for w in descendants(dialog))
            shot = root_dir / "qa/review_20260906/inclusive_purchase.png"
            shot.parent.mkdir(parents=True, exist_ok=True)
            from PIL import ImageGrab
            ImageGrab.grab(bbox=(dialog.winfo_rootx(), dialog.winfo_rooty(), dialog.winfo_rootx()+dialog.winfo_width(), dialog.winfo_rooty()+dialog.winfo_height())).save(shot)
            save = next(w for w in descendants(dialog) if isinstance(w, ttk.Button) and w.cget("text") == "保存修改")
            with patch("pages.purchase_management_page.messagebox.showwarning") as warning:
                save.invoke()
                root.update()
                assert not warning.called, warning.call_args
            assert purchase.get_purchase_order(order)["project_cost_cents"] == 188244
            from pages.import_export_page import ImportExportPage
            from openpyxl import load_workbook
            transfer = ImportExportPage.__new__(ImportExportPage)
            transfer.parent = host
            purchase_file = str(Path(folder) / "purchase.xlsx")
            stored = purchase.get_purchase_order(order)
            with patch("pages.import_export_page.filedialog.asksaveasfilename", return_value=purchase_file), patch("pages.import_export_page.messagebox.showinfo"), patch.object(purchase, "list_purchase_orders", return_value=[stored]):
                transfer.export_purchases()
            book = load_workbook(purchase_file)
            book.active.cell(2, 2, "INCLUSIVE-ROUNDTRIP")
            book.save(purchase_file)
            book.close()
            with patch("pages.import_export_page.filedialog.askopenfilename", return_value=purchase_file), patch("pages.import_export_page.messagebox.showinfo") as info:
                transfer.import_purchases()
                assert info.called, "Import did not succeed"
            imported = purchase.list_purchase_orders(keyword="INCLUSIVE-ROUNDTRIP")
            assert len(imported) == 1 and imported[0]["price_basis"] == "inclusive" and imported[0]["project_cost_cents"] == 188244
            offer_file = str(Path(folder) / "offer.xlsx")
            with patch("pages.import_export_page.filedialog.asksaveasfilename", return_value=offer_file), patch("pages.import_export_page.messagebox.showinfo"), patch.object(master, "list_supplier_offers", return_value=[master.get_supplier_offer(offer) | {"id": offer, "supplier_name": "含税界面测试供应商"}]):
                transfer.export_products()
            with patch("pages.import_export_page.filedialog.askopenfilename", return_value=offer_file), patch("pages.import_export_page.messagebox.showinfo"):
                transfer.import_products()
            assert all(r["price_basis"] == "inclusive" and r["quoted_price_minor"] == 189 for r in master.list_supplier_offers(supplier_id=supplier))
            assert not errors, errors
            print("Desktop acceptance passed: offer basis/input, purchase preview 1.89 x 996 = 1882.44, save/reload, purchase and offer Excel roundtrip")
        finally:
            root.destroy()


if __name__ == "__main__":
    main()
