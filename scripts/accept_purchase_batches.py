"""Exercise the native batch editor using only an isolated database."""
import os
import sys
import tempfile
from pathlib import Path
from unittest.mock import patch


def descendants(widget):
    for child in widget.winfo_children():
        yield child
        yield from descendants(child)


def main():
    project_root = Path(__file__).resolve().parents[1]
    sys.path.insert(0, str(project_root))
    with tempfile.TemporaryDirectory() as folder:
        os.environ["SUPPLY_CHAIN_DB_PATH"] = str(Path(folder) / "test.db")
        os.environ["SUPPLY_CHAIN_ATTACHMENTS_PATH"] = str(Path(folder) / "attachments")
        from db.backup import backup_database
        backup_database(project_root / "supplier_data.db", os.environ["SUPPLY_CHAIN_DB_PATH"])
        from db.migration_runner import run_migrations
        run_migrations()
        from services import procurement_service as purchase, project_service
        from pages.purchase_management_page import PurchaseManagementPage
        import ttkbootstrap as ttk
        from ui.theme import configure_design_system
        project = project_service.create_project({"name": "多材料窗口验收"})
        header = {"purchase_type": "零星采购", "project_id": project, "merchant_name_snapshot": "验收材料店", "purchase_date": "2026-09-12", "freight_amount_cents": 500}
        order = purchase.add_purchase_order(header, {"material_name_snapshot": "钢板", "settlement_mode": "total", "settlement_total_cents": 10000, "price_basis": "inclusive"})
        root = ttk.Window(themename="flatly")
        root.geometry("1200x800+20+20")
        configure_design_system(root)
        errors = []
        root.report_callback_exception = lambda *args: errors.append(str(args))
        try:
            frame = ttk.Frame(root)
            frame.pack(fill="both", expand=True)
            page = PurchaseManagementPage(frame)
            page.open_purchase_dialog("零星采购", order)
            root.update()
            dialog = next(w for w in descendants(root) if isinstance(w, ttk.Toplevel))

            def button(text):
                return next(w for w in descendants(dialog) if isinstance(w, ttk.Button) and w.cget("text") == text)

            def field(label_text):
                label = next(w for w in descendants(dialog) if isinstance(w, ttk.Label) and w.cget("text") == label_text)
                grid = label.grid_info()
                return label.master.grid_slaves(row=grid["row"], column=grid["column"] + 1)[0]

            def set_field(label, value):
                widget = field(label)
                widget.setvar(widget.cget("textvariable"), value)
                root.update()

            with patch("pages.purchase_management_page.messagebox.showwarning") as warning:
                button("加入本单 / 继续填材料").invoke()
                root.update()
                set_field("材料名称 *", "螺丝")
                set_field("整批结算金额（不含运费）*", "35.55")
                button("加入本单 / 继续填材料").invoke()
                root.update()
                tree = next(w for w in descendants(dialog) if isinstance(w, ttk.Treeview))
                assert len(tree.get_children()) == 2
                tree.selection_set("1")
                button("修改选中材料").invoke()
                root.update()
                set_field("整批结算金额（不含运费）*", "40.55")
                assert any(isinstance(w, ttk.Label) and w.cget('textvariable') and w.getvar(w.cget('textvariable')) == '¥40.55' for w in descendants(dialog))
                button("加入本单 / 继续填材料").invoke()
                root.update()
                from PIL import ImageGrab
                output = project_root / "qa" / "purchase_batch_20260912"
                output.mkdir(parents=True, exist_ok=True)
                dialog.attributes("-topmost", True)
                dialog.lift()
                dialog.focus_force()
                root.update()
                import time
                time.sleep(0.2)
                ImageGrab.grab(bbox=(dialog.winfo_rootx(), dialog.winfo_rooty(), dialog.winfo_rootx()+dialog.winfo_width(), dialog.winfo_rooty()+dialog.winfo_height())).save(output / "batch-editor.png")
                button("保存修改").invoke()
                root.update()
                assert not warning.called, warning.call_args_list
                saved = purchase.get_purchase_order(order)
                assert saved["item_count"] == 2
                assert saved["project_cost_cents"] == 14555
                assert len(purchase.list_purchase_orders(project_id=project)) == 1
                from pages.import_export_page import ImportExportPage
                from openpyxl import load_workbook
                transfer = ImportExportPage.__new__(ImportExportPage)
                transfer.parent = root
                export_path = str(Path(folder) / "batch.xlsx")
                detail_rows = purchase.list_purchase_orders(project_id=project, detail_rows=True)
                with patch("pages.import_export_page.filedialog.asksaveasfilename", return_value=export_path), patch("pages.import_export_page.messagebox.showinfo"), patch.object(purchase, "list_purchase_orders", return_value=detail_rows):
                    transfer.export_purchases()
                book = load_workbook(export_path)
                for row_index in (2, 3):
                    book.active.cell(row_index, 2, "BATCH-ROUNDTRIP")
                assert book.active.cell(2, 19).value == 5
                assert book.active.cell(3, 19).value == 0
                book.save(export_path)
                book.close()
                with patch("pages.import_export_page.filedialog.askopenfilename", return_value=export_path), patch("pages.import_export_page.messagebox.showinfo") as imported_message:
                    transfer.import_purchases()
                    assert imported_message.called
                imported = purchase.list_purchase_orders(keyword="BATCH-ROUNDTRIP")
                assert len(imported) == 1 and imported[0]["item_count"] == 2
                assert imported[0]["project_cost_cents"] == 14555
                assert not errors, errors
            print("PASS: native add/edit material, save complete order, freight once, one list row")
        finally:
            root.destroy()


if __name__ == "__main__":
    main()
